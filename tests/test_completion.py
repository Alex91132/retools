# -*- coding: utf-8 -*-
"""Подбор компоновки: точный метод против «Поиска решения» книги и полного перебора."""
import itertools
import json
from pathlib import Path

import pytest

from retools.completion import _dp, evaluate, plan, wc_curve
from retools.rigis import Regressions
from retools.rigis_io import load_project

FIX = Path(__file__).parent / "fixtures"


@pytest.fixture(scope="module")
def project():
    case, obj, pal = load_project(FIX / "rigis_4210H_project.json")
    case.regressions = Regressions()
    return case, obj, pal


@pytest.fixture(scope="module")
def meta():
    return json.loads((FIX / "rigis_4210H_expected.json").read_text(encoding="utf-8"))["meta"]


def test_beats_book_solver(project, meta):
    """Книга («Поиск решения»): 46,3 т/сут при пределе 0 % и 329,3 при 5 %."""
    case, obj, pal = project
    assert plan(case, obj, pal, max_wc=0).q_oil > 7 * meta["solver_book"]["wc_limit_0"]["q_oil"]
    assert plan(case, obj, pal, max_wc=5).q_oil > 1.5 * meta["solver_book"]["wc_limit_5_98"]["q_oil"]


@pytest.mark.parametrize("limit", [0, 5, 10, 50])
def test_constraints_hold(project, limit):
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=limit)
    rules = case.rules
    assert len(c.filters) <= 10
    for f in c.filters:
        assert f.length >= rules.min_filter - 1e-9
        assert f.water_cut <= limit + 1e-6
        assert f.top >= rules.pgmc_bottom - 5.3        # концы фильтров — только работающие срезы
    for a, b in zip(c.filters, c.filters[1:]):
        assert b.top - a.bottom >= rules.min_gap - 1e-9


def test_curve_monotone(project):
    case, obj, pal = project
    q = [c.q_oil for c in wc_curve(case, obj, pal, limits=(0, 5, 10, 20, 100))]
    assert all(b >= a - 1e-6 for a, b in zip(q, q[1:]))


def test_reported_totals_consistent(project):
    """Итог подбора совпадает с независимой оценкой той же компоновки."""
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=5)
    e = evaluate(case, obj, pal, [(f.top, f.bottom) for f in c.filters])
    assert e.q_oil == pytest.approx(c.q_oil, rel=1e-9)


def test_isolations_cover_gaps_and_have_packers(project):
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=5)
    assert c.isolations
    for iso in c.isolations:
        assert iso.packers
        for p0, p1 in iso.packers:
            assert p1 - p0 == pytest.approx(case.rules.elastomer)
            assert iso.top - 1e-9 <= p0 and p1 <= iso.bottom + 1e-9


def test_dp_matches_brute_force_on_small_problem():
    """Динамическое программирование = полный перебор (8 срезов, до 3 фильтров)."""
    tops = [0, 5, 10, 15, 20, 25, 30, 35]
    bots = [5, 10, 15, 20, 25, 30, 35, 40]
    ql = [10, 12, 30, 8, 14, 40, 9, 11]
    qo = [9, 11, 6, 7, 13, 5, 8, 10]
    rho, max_wc, min_f, gap, kmax = 1.0, 30.0, 10.0, 5.0, 3
    segs, best = _dp(tops, bots, ql, qo, rho, max_wc, min_f, gap, kmax)

    def ok(seg):
        a, b = seg
        if bots[b] - tops[a] < min_f:
            return False
        sl, so = sum(ql[a:b + 1]), sum(qo[a:b + 1])
        return so > 0 and so >= (1 - max_wc / 100) * rho * sl

    cand = [(a, b) for a in range(8) for b in range(a, 8) if ok((a, b))]
    brute = 0.0
    for k in range(1, kmax + 1):
        for comb in itertools.combinations(cand, k):
            comb = sorted(comb)
            if all(tops[c2[0]] - bots[c1[1]] >= gap for c1, c2 in zip(comb, comb[1:])):
                brute = max(brute, sum(sum(qo[a:b + 1]) for a, b in comb))
    assert best == pytest.approx(brute)


def test_liner_is_contiguous_and_typed(project):
    from retools.completion import liner
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=5)
    els = liner(c, case.rules)
    assert els[0].kind == "фильтр" and els[0].top == pytest.approx(c.filters[0].top)
    for a, b in zip(els, els[1:]):
        assert b.top == pytest.approx(a.bottom)
    assert {e.kind for e in els} <= {"фильтр", "глухая труба", "эластомер"}
    assert sum(e.length for e in els if e.kind == "фильтр") == pytest.approx(sum(f.length for f in c.filters))


def test_short_gap_gets_single_centered_packer(project):
    """Промежуток короче 2·(отступ + эластомер): два пакера не помещаются — один посередине."""
    from dataclasses import replace
    case, obj, pal = project
    tight = replace(case, rules=replace(case.rules, elastomer_offset=10.0))
    c = plan(tight, obj, pal, max_wc=5)
    for iso in c.isolations:
        gap = iso.bottom - iso.top
        if gap < 2 * (10.0 + tight.rules.elastomer) and len(iso.packers) == 1 and iso.top > c.filters[0].top:
            p0, p1 = iso.packers[0]
            assert (p0 + p1) / 2 == pytest.approx((iso.top + iso.bottom) / 2)


def test_excel_report(project):
    import io
    import openpyxl
    from retools.report import completion_xlsx
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=5)
    wb = openpyxl.load_workbook(io.BytesIO(completion_xlsx(case, c, [c])))
    assert wb.sheetnames == ["Компоновка", "Изоляция", "Срезы", "Кривая", "Настройки"]
    assert wb["Срезы"].max_row == len(case.slices) + 1


@pytest.mark.parametrize("limit", [5, 8])
def test_well_level_water_cut(project, limit):
    """Предел по скважине: суммарная обводнённость ≤ предела, Qн ≥ чем при пределе по каждому фильтру."""
    case, obj, pal = project
    well = plan(case, obj, pal, max_wc=limit, wc_mode="well")
    per_filter = plan(case, obj, pal, max_wc=limit, wc_mode="filter")
    assert well.water_cut <= limit + 1e-6
    assert well.q_oil >= per_filter.q_oil - 1e-6


def test_gas_offset_keeps_filters_away_from_gas(project):
    """Синтетика: срезы 3300–3320 объявлены газом — с отступом 30 м фильтры туда не заходят."""
    from dataclasses import replace
    case, obj, pal = project
    sl = [replace(s, saturation="газ") if 3300 <= s.top < 3320 else s for s in case.slices]
    gcase = replace(case, slices=sl)
    c = plan(gcase, obj, pal, max_wc=20, gas_offset=30)
    for f in c.filters:
        assert f.bottom <= 3300 - 30 + 6 or f.top >= 3320 + 30 - 6     # с точностью до среза
    assert any(i.reason in ("газ", "отступ от газа") for i in c.isolations)


# --- МЗС: несколько стволов ---------------------------------------------------

def test_plan_multi_single_lateral_equals_plan(project):
    from retools.completion import plan_multi
    case, obj, pal = project
    a = plan(case, obj, pal, max_wc=5)
    b = plan_multi([case], obj, pal, max_wc=5)
    assert b.q_oil == pytest.approx(a.q_oil, rel=1e-12)
    assert [(f.top, f.bottom) for f in b.laterals[0].filters] == [(f.top, f.bottom) for f in a.filters]


def test_two_laterals_share_filter_budget_and_section_length(project):
    """Два одинаковых ствола: общий лимит фильтров соблюдается, суммарная длина секций
    больше — дебит каждого ствола ниже, чем у одиночной ГС (коэффициент секции)."""
    from dataclasses import replace
    from retools.completion import plan_multi
    case, obj, pal = project
    lat2 = replace(case, slices=[replace(s, well="4210H/2") for s in case.slices])
    single = plan(case, obj, pal, max_wc=5)
    w = plan_multi([case, lat2], obj, pal, max_wc=5, max_filters=4)
    assert sum(len(c.filters) for c in w.laterals) <= 4
    assert w.l_total > single.l_total
    assert w.q_oil > single.q_oil
    assert max(c.q_oil for c in w.laterals) < single.q_oil + 1e-6
    assert w.names == ["4210H", "4210H/2"]


def test_combine_matches_brute_force():
    """Распределение фильтров между двумя стволами = полный перебор."""
    from retools.completion import _combine, _dp_table
    rng = __import__("random").Random(7)
    lats = []
    for _ in range(2):
        n = 7
        tops = [5.0 * i for i in range(n)]
        bots = [5.0 * (i + 1) for i in range(n)]
        ql = [rng.uniform(5, 30) for _ in range(n)]
        qo = [q * rng.uniform(0.5, 1.0) for q in ql]
        lats.append((tops, bots, ql, qo))
    K = 3
    tables = [_dp_table(t, b, ql, qo, 1.0, 30.0, 10.0, 5.0, K) for t, b, ql, qo in lats]
    segs = _combine(tables, K)
    got = sum(sum(qo[a:b + 1]) for (t, bb, ql, qo), sg in zip(lats, segs) for a, b in sg)
    best = max(tables[0][0][k1] + tables[1][0][k2] for k1 in range(K + 1) for k2 in range(K + 1 - k1)
               if tables[0][0][k1] > -1e17 and tables[1][0][k2] > -1e17)
    assert got == pytest.approx(best)


def test_filter_cost_drops_marginal_filters(project):
    """С минимальным приростом 5 т/сут каждый фильтр даёт ≥ 5 т/сут, фильтров не больше."""
    case, obj, pal = project
    free = plan(case, obj, pal, max_wc=0)
    costly = plan(case, obj, pal, max_wc=0, filter_cost=5.0)
    assert len(costly.filters) <= len(free.filters)
    assert all(f.q_oil >= 5.0 - 1e-6 for f in costly.filters)
    assert costly.q_oil <= free.q_oil + 1e-6


def test_owc_offset_by_tvd(project):
    """ВНК на АО 1611 с отступом 3 м: вскрываются только срезы с подошвой выше 1608 м."""
    case, obj, pal = project
    c = plan(case, obj, pal, max_wc=20, owc_tvd=1611.0, owc_offset=3.0)
    for f in c.filters:
        for s in case.slices:
            if f.top <= s.top and s.bottom <= f.bottom and s.bottom_tvd > 0:
                assert s.bottom_tvd <= 1608.0 + 1e-9
    assert any(i.reason == "отступ от ВНК" for i in c.isolations)


def test_forbidden_mask_reasons():
    from retools.completion import forbidden_mask
    from retools.rigis import Slice
    sl = [Slice("w", "p", 0, 5, 1000, 1001, saturation="газ"), Slice("w", "p", 5, 10, 1001, 1002),
          Slice("w", "p", 10, 15, 1010, 1011), Slice("w", "p", 15, 20, 1019, 1020)]
    m = forbidden_mask(sl, gas_offset=3, goc_tvd=None, owc_tvd=1021, owc_offset=2)
    assert m == ["отступ от газа", "отступ от газа", "", "отступ от ВНК"]


def test_manual_plan_reproduces_auto_completion():
    """Интервалы подобранной компоновки, введённые вручную, дают те же дебиты, изоляцию и пакеры."""
    from retools.completion import manual_plan, plan_multi
    case, obj, pal = load_project(FIX / "rigis_4210H_project.json")
    auto = plan_multi([case], obj, pal, max_wc=5)
    ranges = [[(f.top, f.bottom) for f in auto.laterals[0].filters]]
    man, notes = manual_plan([case], obj, pal, ranges, max_wc=5)
    assert man.q_oil == pytest.approx(auto.q_oil, rel=1e-12) and man.q_liq == pytest.approx(auto.q_liq, rel=1e-12)
    assert [(i.top, i.bottom, i.packers) for i in man.laterals[0].isolations] == \
        [(i.top, i.bottom, i.packers) for i in auto.laterals[0].isolations]
    assert notes == []


def test_manual_plan_reports_rule_violations():
    from retools.completion import manual_plan
    case, obj, pal = load_project(FIX / "rigis_4210H_project.json")
    s = [x for x in case.slices if x.collector.strip() == "К"]
    # фильтр 1 — один срез (4,3 м < 10), промежуток — один срез (6,3 м < 15), фильтр 2 — два среза
    ranges = [[(s[100].top, s[100].bottom), (s[102].top, s[103].bottom)]]
    man, notes = manual_plan([case], obj, pal, ranges)
    assert any("короче" in n for n in notes) and any("промежуток" in n for n in notes)
    assert man.q_oil > 0
