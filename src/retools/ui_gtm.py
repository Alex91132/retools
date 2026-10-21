# -*- coding: utf-8 -*-
"""Страница Streamlit «Расчёт добычи по рейтингу ГТМ» (порт листов «Рейтинг ГТМ» и «Расчёт добычи»).

Сценарий: загрузить книгу RE Tools (рейтинг и листы объектов) или пример →
поправить параметры скважин в таблице → «Рассчитать» → запасы, КИН и
стартовые дебиты по скважинам, профиль добычи по месяцам (сумма и по
скважине) → выгрузка в Excel.
"""
from __future__ import annotations

import datetime as dt
import io
import json
import tempfile
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

from .gtm_calc import ObjectSheet, production_rows, rating_row, result_by_month

FIX = Path(__file__).resolve().parents[2] / "tests" / "fixtures"

# столбцы рейтинга, которые редактируются в таблице: (буква, заголовок)
RATING_FIELDS = [
    ("A", "№"), ("E", "Скважина"), ("F", "Объект"), ("G", "Лист объекта"), ("H", "ГТМ"), ("I", "Вариант"),
    ("J", "Дата запуска"), ("M", "h эфф, м"), ("N", "Кп"), ("O", "k, мД"), ("P", "Kv/Kh"), ("Q", "Sw нач."),
    ("R", "Рпл нач., атм"), ("S", "Доля заводн. толщин"), ("T", "L ГС общ., м"), ("U", "L ГС эфф., м"),
    ("V", "Число БС"), ("W", "L БС общ., м"), ("X", "L БС эфф., м"), ("Y", "Число ГРП"), ("AA", "Xf, м"),
    ("AB", "Wf, мм"), ("AC", "k проппанта, Д"), ("AE", "Трещины"), ("AF", "Rконт, м"), ("AG", "Rскв, м"),
    ("AH", "Скин"), ("AI", "Рзаб, атм"), ("AJ", "Tзап, сут"), ("AL", "Начало ППД, мес"), ("AM", "ППД/доб"),
    ("AN", "Рбуф ППД, атм"), ("AP", "Множ. Jd ППД (авто-ГРП)"), ("AQ", "Макс. приёмист., м³/сут"),
    ("AR", "Макс. компенсация"), ("AS", "Аквифер, V/Vпл"), ("AT", "Мин. Qн, т/сут"), ("AU", "Пред. обв., %"),
    ("AV", "Макс. депрессия, атм"), ("AW", "Коэф. успешн. Jd"), ("AX", "Кэ"), ("AY", "Длительность, мес"),
]
TEXT_COLS = {"E", "F", "G", "H", "I", "AE"}


def _example():
    book = json.loads((FIX / "production_book.json").read_text(encoding="utf-8"))
    from .gtm_io import RATING_INPUT_COLS
    rows = [{k: r.get(k) for k in RATING_INPUT_COLS if r.get(k) is not None} for r in book["rating"]]
    obj = ObjectSheet.from_book(json.loads((FIX / "pvt_mbe_book.json").read_text(encoding="utf-8")))
    return rows, {obj.name: obj}


def _from_book(upload):
    from .gtm_io import _load_values, read_object_sheet, read_rating_rows
    f = tempfile.NamedTemporaryFile(delete=False, suffix=Path(upload.name).suffix)
    f.write(upload.getbuffer())
    f.close()
    wb = _load_values(f.name)
    rows = read_rating_rows(wb=wb)
    objs = {name: read_object_sheet(wb=wb, sheet=name) for name in {r.get("G") for r in rows if r.get("G")}}
    return rows, objs


def _to_df(rows: list[dict]) -> pd.DataFrame:
    data = []
    for r in rows:
        d = {}
        for c, h in RATING_FIELDS:
            v = r.get(c)
            if c == "J" and v is not None and not isinstance(v, (dt.date, dt.datetime)):
                v = dt.datetime.fromisoformat(str(v))
            if c == "J" and isinstance(v, dt.datetime):
                v = v.date()
            d[h] = v if v is not None else ("" if c in TEXT_COLS else 0.0)
        data.append(d)
    return pd.DataFrame(data)


def _from_df(df: pd.DataFrame, base: list[dict]) -> list[dict]:
    out = []
    for i, row in df.reset_index(drop=True).iterrows():
        r = dict(base[i]) if i < len(base) else {}
        for c, h in RATING_FIELDS:
            v = row[h]
            if c in TEXT_COLS:
                if isinstance(v, float) and not pd.isna(v) and v.is_integer():
                    v = int(v)                      # номер скважины из таблицы приходит числом
                r[c] = None if v in ("", None) or (isinstance(v, float) and pd.isna(v)) else str(v)
            elif c == "A":
                r[c] = None if pd.isna(v) else (int(v) if float(v).is_integer() else float(v))
            elif c == "J":
                r[c] = v if isinstance(v, dt.date) else dt.date.today()
            else:
                r[c] = None if pd.isna(v) else float(v)
        r["T"] = r.get("T") or 0
        out.append(r)
    return out


def _xlsx(rating: list[dict], rows: list[dict], months: list[dict]) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    def sheet(ws, header, data):
        ws.append(header)
        for c in ws[ws.max_row]:
            c.font = bold
        for d in data:
            ws.append(d)

    ws = wb.active
    ws.title = "Рейтинг"
    cols = [("A", "№"), ("E", "Скважина"), ("K", "Тип"), ("BC", "Qж, м³/сут"), ("BE", "Qн, т/сут"), ("BD", "Обв., %"),
            ("BF", "ГФ, м³/т"), ("DF", "ПИ, м³/сут/атм"), ("CF", "Jd"), ("AZ", "НГЗ, тыс. т"), ("BB", "НИЗ, тыс. т"),
            ("BA", "КИН потенц.")]
    sheet(ws, [h for _, h in cols], [[r.get(c) for c, _ in cols] for r in rating])
    ws = wb.create_sheet("Расчёт добычи")
    cols = [("A", "№"), ("E", "Скважина"), ("FL", "Дата"), ("FM", "Месяц"), ("FN", "Дни"), ("BC", "Qж, м³/сут"),
            ("BE", "Qн, т/сут"), ("BD", "Обв., %"), ("BF", "ГФ, м³/т"), ("CL", "Рпл истощ., атм"),
            ("DY", "Рпл завод., атм"), ("GJ", "Рпл средн., атм"), ("GK", "Рзаб средн., атм"), ("DT", "Закачка, м³/сут"),
            ("FO", "Нефть, тыс. т"), ("FS", "Жидкость, тыс. м³"), ("FV", "Нак. нефть, тыс. т"),
            ("FZ", "Нак. жидкость, тыс. м³"), ("GA", "Нак. закачка, тыс. м³"), ("GS", "КИН")]
    sheet(ws, [h for _, h in cols], [[r.get(c) for c, _ in cols] for r in rows])
    ws = wb.create_sheet("По месяцам")
    keys = [("date", "Дата"), ("n_wells", "Скважин"), ("q_oil", "Qн, т/сут"), ("q_liq", "Qж, м³/сут"),
            ("water_cut", "Обв., %"), ("inj", "Закачка, м³/сут"), ("oil_t", "Нефть, тыс. т"),
            ("liq_m3", "Жидкость, тыс. м³"), ("cum_oil_t", "Нак. нефть, тыс. т"), ("cum_liq_m3", "Нак. жидкость, тыс. м³"),
            ("cum_inj", "Нак. закачка, тыс. м³"), ("p_res", "Рпл, атм"), ("p_wf", "Рзаб, атм")]
    sheet(ws, [h for _, h in keys], [[m[k] for k, _ in keys] for m in months])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render():
    st.header("Расчёт добычи по рейтингу ГТМ")
    st.caption("Порт листов «Рейтинг ГТМ» и «Расчёт добычи» RE Tools: запасы и потенциальный КИН, продуктивность "
               "Get_J (ННС, ГРП, ГС, МГРП, МЗС), помесячный расчёт истощаемой и заводняемой частей пласта с ППД. "
               "Сверено с книгой: 72 строки расчёта совпадают до 1e-10 при той же продуктивности.")
    state = st.session_state
    src = st.radio("Источник", ["Пример книги (13 скважин)", "Книга RE Tools (.xlsm)"], horizontal=True, key="gtm_src")
    if src.startswith("Книга"):
        up = st.file_uploader("Книга RE Tools", type=["xlsm", "xlsx"], key="gtm_book")
        if up is not None and state.get("gtm_loaded") != up.name:
            with st.spinner("Читаю рейтинг и листы объектов (макросы не выполняются)…"):
                state.gtm_rows, state.gtm_objs = _from_book(up)
                state.gtm_loaded = up.name
    elif state.get("gtm_loaded") != "example":
        state.gtm_rows, state.gtm_objs = _example()
        state.gtm_loaded = "example"
    if "gtm_rows" not in state:
        st.info("Загрузите книгу.")
        return
    objs = state.gtm_objs
    c1, c2 = st.columns(2)
    rebuild = c1.checkbox("Объект пересчитать из параметров (матбаланс и палетки заново)", value=False,
                          help="Иначе — таблицы и палетки листа объекта из книги. В книге ниже ~124 атм матбаланс не "
                               "досчитан, поэтому при низком Рпл результаты будут отличаться.")
    cj_book = c2.checkbox("CJ как в книге для строк 8+ (НПЗ с приростом от ППД)", value=False,
                          help="В строках 8 и ниже листа «Рейтинг ГТМ» НПЗ вытеснения включает прирост от ППД, "
                               "который расчёт добычи прибавляет ещё раз. По умолчанию — формула строки 7.")
    if rebuild:
        key = ("rebuilt", tuple(sorted(objs)))
        if state.get("gtm_objs_key") != key:
            with st.spinner("Решаю матбаланс и строю палетки…"):
                state.gtm_objs_built = {n: ObjectSheet.from_params(o.params, name=n, piezo_lag=o.piezo_lag,
                                                                    ma_period=o.ma_period) for n, o in objs.items()}
                state.gtm_objs_key = key
        objs = state.gtm_objs_built
    st.subheader("Скважины")
    df = st.data_editor(_to_df(state.gtm_rows), num_rows="dynamic", use_container_width=True, key="gtm_table",
                        hide_index=True)
    rows_in = _from_df(df, state.gtm_rows)
    missing = sorted({r.get("G") for r in rows_in} - set(objs))
    if missing:
        st.error("Нет листа объекта: " + ", ".join(map(str, missing)))
        return
    key = (repr(rows_in), rebuild, cj_book)
    if st.button("Рассчитать", type="primary") or state.get("gtm_key") != key:
        bar = st.progress(0.0, text="Рейтинг…")
        rating = [rating_row(r, objs[r["G"]], cj_with_ci=cj_book and i > 0) for i, r in enumerate(rows_in)]

        def prog(i, total, row):
            bar.progress(min(i / max(total, 1), 1.0), text=f"Скважина {min(i + 1, total)} из {total}, "
                                                           f"месяц {row['FM']}")

        rows = production_rows(rating, objs, progress=prog)
        bar.empty()
        state.gtm_res = (rating, rows, result_by_month(rows))
        state.gtm_key = key
    rating, rows, months = state.gtm_res
    st.subheader("Рейтинг: запасы и стартовые параметры")
    st.dataframe(pd.DataFrame([{
        "№": r["A"], "Скважина": r["E"], "Тип": r["K"], "Qж, м³/сут": round(r["BC"], 1), "Qн, т/сут": round(r["BE"], 1),
        "Обв., %": round(r["BD"], 1), "ПИ": round(r["DF"], 3), "Jd": round(r["CF"], 3),
        "НГЗ, тыс. т": round(r["AZ"], 1), "НИЗ, тыс. т": round(r["BB"], 1), "КИН потенц.": round(r["BA"], 3)}
        for r in rating]), hide_index=True, use_container_width=True)
    st.subheader("Профиль добычи")
    dates = [m["date"] for m in months]
    fig = make_subplots(rows=1, cols=3, subplot_titles=("Qн, т/сут и Qж, м³/сут", "Обводнённость, %", "Рпл и Рзаб, атм"))
    fig.add_trace(go.Scatter(x=dates, y=[m["q_oil"] for m in months], name="Qн", line=dict(color="#2e9d4f")), 1, 1)
    fig.add_trace(go.Scatter(x=dates, y=[m["q_liq"] for m in months], name="Qж", line=dict(color="#2f7fd8")), 1, 1)
    fig.add_trace(go.Scatter(x=dates, y=[m["water_cut"] for m in months], name="обв.", line=dict(color="#d64545"),
                             showlegend=False), 1, 2)
    fig.add_trace(go.Scatter(x=dates, y=[m["p_res"] for m in months], name="Рпл", line=dict(color="#8e44ad")), 1, 3)
    fig.add_trace(go.Scatter(x=dates, y=[m["p_wf"] for m in months], name="Рзаб", line=dict(color="#9aa0a6")), 1, 3)
    fig.update_layout(height=330, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.25))
    st.plotly_chart(fig, use_container_width=True)
    by_year = {}
    for m in months:
        y = by_year.setdefault(m["date"].year, {"Год": m["date"].year, "Нефть, тыс. т": 0.0, "Жидкость, тыс. м³": 0.0})
        y["Нефть, тыс. т"] += m["oil_t"]
        y["Жидкость, тыс. м³"] += m["liq_m3"]
        y["Нак. нефть, тыс. т"] = m["cum_oil_t"]
    st.dataframe(pd.DataFrame([{k: (round(v, 2) if isinstance(v, float) else v) for k, v in y.items()}
                               for y in by_year.values()]), hide_index=True)
    names = [f"{r['A']} · {r['E']} ({r['K']})" for r in rating]
    pick = st.selectbox("Скважина", names)
    a = rating[names.index(pick)]["A"]
    w = [r for r in rows if r["A"] == a]
    if w:
        fig = make_subplots(rows=1, cols=3, subplot_titles=("Qн / Qж", "Обводнённость, %", "Рпл истощ. / завод., атм"))
        t = [r["FL"] for r in w]
        fig.add_trace(go.Scatter(x=t, y=[r["BE"] for r in w], name="Qн", line=dict(color="#2e9d4f")), 1, 1)
        fig.add_trace(go.Scatter(x=t, y=[r["BC"] for r in w], name="Qж", line=dict(color="#2f7fd8")), 1, 1)
        fig.add_trace(go.Scatter(x=t, y=[r["BD"] for r in w], name="обв.", line=dict(color="#d64545"),
                                 showlegend=False), 1, 2)
        fig.add_trace(go.Scatter(x=t, y=[r["CL"] for r in w], name="Рпл истощ.", line=dict(color="#8e44ad")), 1, 3)
        fig.add_trace(go.Scatter(x=t, y=[r["DY"] for r in w], name="Рпл завод.", line=dict(color="#e07b00")), 1, 3)
        fig.update_layout(height=300, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.3))
        st.plotly_chart(fig, use_container_width=True)
    st.download_button("Расчёт в Excel", _xlsx(rating, rows, months), file_name="расчёт_добычи.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
