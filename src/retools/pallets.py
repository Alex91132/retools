# -*- coding: utf-8 -*-
"""Палетки заводнения листа «PVT_MBE_Waterflood» и билинейная интерполяция.

Три палетки книги (строки 69-84):
- **Mu** — вязкость смеси от (Рпл × Sw): строки 70-84 = давления PVT-таблицы
  (410,5 … 159 атм), колонки = водонасыщенности (0,772 … 0,2);
- **Kro Mix** и **Krw Mix** — ОФП смеси нефть-вода от (%НПЗ × Sw), где %НПЗ —
  доля отобранных извлекаемых запасов (0 … 1 с шагом по сетке книги).

Интерполяция — билинейная (порт поведения crv_interpolation_2D с типом 0 →
alglib Spline2DBuildBilinear/Spline2DCalc): в узлах сетки значения точны,
между узлами — плоскость по ячейке, вне сетки — продление крайними ячейками
(ограничение по границам, безопаснее экстраполяции alglib для палеток).
"""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import openpyxl

BOOK_SHEET = "PVT_MBE_Waterflood"
PALLET_COLS = {"mu": 3, "kro_mix": 17, "krw_mix": 31}   # C, Q, AE
N_ROWS, N_COLS = 15, 12


@dataclass(frozen=True)
class Grid:
    """Регулярная сетка: оси x (колонки), y (строки), значения f[y][x]."""
    x: tuple
    y: tuple
    f: tuple

    def value(self, x: float, y: float) -> float:
        """Билинейная интерполяция по сетке (вне сетки — по краевой ячейке).

        Оси палеток книги убывающие (P: 410→159, Sw: 0,772→0,2) — locate
        работает с любым направлением оси.
        """
        def locate(axis, v):
            n = len(axis)
            asc = axis[1] > axis[0]
            for i in range(n - 1):
                a, b = axis[i], axis[i + 1]
                lo, hi = (a, b) if asc else (b, a)
                if lo <= v <= hi:
                    t = (v - a) / (b - a) if b != a else 0.0
                    return i, min(max(t, 0.0), 1.0)
            beyond_max = v > max(axis)
            if asc:
                return (n - 2, 1.0) if beyond_max else (0, 0.0)
            return (0, 0.0) if beyond_max else (n - 2, 1.0)

        ix, tx = locate(self.x, x)
        iy, ty = locate(self.y, y)
        f00 = self.f[iy][ix]
        f10 = self.f[iy][ix + 1]
        f01 = self.f[iy + 1][ix]
        f11 = self.f[iy + 1][ix + 1]
        return (f00 * (1 - tx) * (1 - ty) + f10 * tx * (1 - ty)
                + f01 * (1 - tx) * ty + f11 * tx * ty)


def read_pallets(book_path: str | Path) -> dict[str, Grid]:
    """Прочитать три палетки из книги RE Tools (кэшированные значения)."""
    wb = openpyxl.load_workbook(book_path, read_only=True, data_only=True)
    ws = wb[BOOK_SHEET]
    out = {}
    for name, col0 in PALLET_COLS.items():
        y = tuple(ws.cell(row=r, column=col0).value for r in range(70, 70 + N_ROWS))
        x = tuple(ws.cell(row=69, column=col0 + 1 + j).value for j in range(N_COLS))
        f = tuple(tuple(ws.cell(row=r, column=col0 + 1 + j).value for j in range(N_COLS))
                  for r in range(70, 70 + N_ROWS))
        out[name] = Grid(x=x, y=y, f=f)
    wb.close()
    return out


def default_pallets() -> dict[str, Grid]:
    """Палетки контрольного примера книги (из фикстуры тестов)."""
    data = json.loads((Path(__file__).resolve().parents[2] / "tests"
                       / "pallets_fixture.json").read_text(encoding="utf-8"))
    return {k: Grid(x=tuple(v["x"]), y=tuple(v["y"]),
                    f=tuple(tuple(row) for row in v["f"]))
            for k, v in data.items()}


def mu_mix(pallets: dict[str, Grid], p: float, sw: float) -> float:
    """Вязкость смеси по палетке Mu (аналог вызова книги в РИГИС: ES1)."""
    return pallets["mu"].value(sw, p)


def kro_mix(pallets: dict[str, Grid], npz_fraction: float, sw: float) -> float:
    """ОФП нефти смеси по палетке Kro Mix (%НПЗ = доля извлечённых запасов)."""
    return pallets["kro_mix"].value(sw, npz_fraction)


def krw_mix(pallets: dict[str, Grid], npz_fraction: float, sw: float) -> float:
    """ОФП воды смеси по палетке Krw Mix."""
    return pallets["krw_mix"].value(sw, npz_fraction)
