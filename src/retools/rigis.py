# -*- coding: utf-8 -*-
"""Расчёт интервалов РИГИС по срезам — порт формул листа «РИГИС» (строка 70, R:CE).

Вход: срезы РИГИС (глубины, мощность, УЭС, Кп, Кгл, Кпр, Кв, Кн, характер
насыщения, коллектор), настройки листа (блок «Выбор» — интервалы и параметры
пласта по интервалам, параметры по умолчанию, регрессии петрофизики, правила
компоновки), объект разработки (PVT, ОФП, матбаланс) и палетки секций ГС.

Каждый срез ГС считается как секция горизонтального ствола длиной H в пласте
эффективной мощностью h_eff (Get_J_obsolete), с поправкой на взаимовлияние
секций CA(L общ. эфф., L секции, h_eff) из палеток нормировки Jd.
Обозначения полей — по шапке листа (строка 72), буква колонки — в комментарии.
"""
from __future__ import annotations

import math
from dataclasses import dataclass, field, replace
from typing import Iterable, Sequence

from . import rigis_physics as ph
from .field_object import FieldObject, JdSectionPallets


@dataclass(frozen=True)
class Slice:
    """Срез РИГИС (колонки B:Q листа)."""
    well: str
    formation: str
    top: float               # D кровля, м (MD)
    bottom: float            # E подошва, м (MD)
    top_tvd: float = 0.0     # F
    bottom_tvd: float = 0.0  # G
    h: float = 0.0           # H мощность (длина среза по стволу), м
    res: float = 0.0         # I УЭС, Омм
    phit: float = 0.0        # J Кп общ. (PHIT)
    phi: float = 0.0         # K Кп для расчёта
    vsh: float = 0.0         # L Кгл
    perm: float = 0.0        # M Кпр, мД
    sw: float = 0.0          # N Кв
    so: float = 0.0          # O Кн
    saturation: str = ""     # P характер насыщения
    collector: str = ""      # Q коллектор («К»)

    @property
    def is_gas(self) -> bool:
        return self.saturation.strip().lower() == "газ"


@dataclass(frozen=True)
class IntervalParams:
    """Строка блока «Выбор» (B:K, строки 4-13): интервал и параметры пласта."""
    n: int
    top: float = 0.0
    bottom: float = 0.0
    h_eff: float = 0.0
    p_res: float = 0.0
    flood_share: float = 0.0
    r_contour: float = 0.0
    inj_ratio: float = 0.0


@dataclass(frozen=True)
class Defaults:
    """Строка 15: параметры скважины и пласта по умолчанию."""
    kv_kh: float = 0.1       # C
    r_w: float = 0.1         # D
    t_pss: float = 5.0       # E, сут
    p_wf: float = 82.0       # F целевое Рзаб
    h_eff: float = 15.0      # G
    p_res: float = 110.0     # H
    flood_share: float = 0.7  # I
    r_contour: float = 330.0  # J
    inj_ratio: float = 0.8   # K


@dataclass(frozen=True)
class Regressions:
    """Петрофизика объекта (AK:AQ, строка 5): Кво, Кно, Квкрит, УЭС граничное.

    a = 0 означает «брать константу объекта»: Swl, Sorw_ow, Swcr_ow.
    """
    kno_a: float = 0.0       # AK: Кно = (1-Кво)(1-(a·ln(Кпр)+b))
    kno_b: float = 0.0       # AL
    kvo_a: float = 0.0       # AM: Кво = -a·Кп + b
    kvo_b: float = 0.0       # AN
    kvcr_a: float = 0.0      # AO: Квкрит = -a·Кп·100 + b
    kvcr_b: float = 0.0      # AP
    res_cut: float = 0.0     # AQ УЭС граничное


@dataclass(frozen=True)
class CompletionRules:
    """Правила компоновки (S4:T14): ПГМЦ, длины фильтров и пакеров."""
    pgmc_top: float = 0.0           # S4 — интервал ПГМЦ (выше подошвы срезы не работают)
    pgmc_bottom: float = 0.0        # T4
    min_filter: float = 10.0        # T5 мин. длина фильтра
    elastomer_offset: float = 5.0   # T7 расстояние до эластомера
    min_gap: float = 15.0           # T9 мин. расстояние между фильтрами
    boundary: float = 50.0          # T11 «Граница»
    blank_pipe: float = 11.0        # T12 глухая труба
    elastomer: float = 4.5          # T13 длина эластомера
    max_wc: float = 5.0             # T14 макс. % обв. для автоподбора


@dataclass
class RigisCase:
    """Всё, что нужно для расчёта скважины."""
    slices: list
    intervals: list                 # IntervalParams (до 10)
    defaults: Defaults
    regressions: Regressions
    rules: CompletionRules
    k_success: float = 1.0          # AF18 Кусп
    object_name: str = "PVT_MBE_Waterflood"
    well_type: str = "horizontal"   # "horizontal" (ГС/МЗС — логика книги) | "vertical" (перфорация)


@dataclass
class SliceResult:
    """Результат расчёта среза (колонки R:CE)."""
    interval: int = 0       # R / T
    used: bool = False      # S
    kvo: float = 0.0        # AA
    kno: float = 0.0        # AB
    kvcr: float = 0.0       # AC
    q_liq: float = 0.0      # AE, м³/сут
    fw: float = 0.0         # AF, %
    q_oil: float = 0.0      # AG, т/сут
    gor: float = 0.0        # AH
    oip: float = 0.0        # AI НГЗ, тыс. т
    rec: float = 0.0        # AJ извлекаемые, тыс. т
    h_eff: float = 0.0      # AL
    p_res: float = 0.0      # AM
    r_contour: float = 0.0  # AO
    flood_share: float = 0.0  # AN
    jd: float = 0.0         # AQ
    sg: float = 0.0         # BI
    so: float = 0.0         # BJ
    kro: float = 0.0        # BK
    krg: float = 0.0        # BL
    krw: float = 0.0        # BM
    mu: float = 0.0         # BN
    ct: float = 0.0         # BO
    b: float = 0.0          # BP
    pi: float = 0.0         # BQ
    dp: float = 0.0         # BR
    l_total: float = 0.0    # BS
    l_section: float = 0.0  # BT
    k_section: float = 0.0  # CA
    pvt: dict = field(default_factory=dict)


def assign_intervals(slices: Sequence[Slice], ranges: Sequence[tuple]) -> list[int]:
    """Колонка R: номер интервала (1..n), в который попадает кровля среза, или 0."""
    out = []
    for s in slices:
        n = 0
        for i, (t, b) in enumerate(ranges, start=1):
            if t <= s.top < b:
                n = i
                break
        out.append(n)
    return out


def _interval_value(intervals: Sequence[IntervalParams], n: int, attr: str, default: float) -> float:
    if n == 0:
        return default
    for iv in intervals:
        if iv.n == n:
            v = getattr(iv, attr)
            return default if not v else v
    return default


def compute(case: RigisCase, obj: FieldObject, pallets: JdSectionPallets,
            ranges: Sequence[tuple] | None = None, l_total: float | None = None,
            p_res_override: Sequence[float] | None = None) -> list[SliceResult]:
    """Расчёт всех срезов при заданных интервалах компоновки.

    ranges — интервалы (кровля, подошва) блока «Выбор» (C:D); по умолчанию — из case.
    Они задают и номер интервала (параметры пласта по интервалам), и компоновку
    для длины BS — как в книге.
    l_total — суммарная длина фильтров для коэффициента секции (BS) у всех
    работающих срезов; нужна подбору компоновки, где фильтры не совпадают с
    интервалами параметров. None — как в книге (по «Компкол»).
    p_res_override — текущее пластовое давление по срезам (для прогноза во времени).
    """
    sl = case.slices
    if ranges is None:
        ranges = [(iv.top, iv.bottom) for iv in case.intervals]
    d, rg, rules, prm = case.defaults, case.regressions, case.rules, obj.params
    rho = obj.rho_oil
    r_col = assign_intervals(sl, ranges)
    used = [s.collector.strip() == "К" and s.sw > 0 and s.bottom > rules.pgmc_bottom for s in sl]
    t_col = [r if u else 0 for r, u in zip(r_col, used)]
    # BS: длина всех срезов в компоновке (Компкол > 0) или вне её (= 0); BU: длина по интервалу
    len_in = sum(s.h for s, t in zip(sl, t_col) if t > 0)
    len_out = sum(s.h for s, t in zip(sl, t_col) if t == 0)
    by_interval: dict[int, float] = {}
    for s, t in zip(sl, t_col):
        by_interval[t] = by_interval.get(t, 0.0) + s.h

    t_C = prm["t_C"]
    sg_pwf = obj.sg(d.p_wf)
    bo_pwf = obj.bo_curve_at(d.p_wf)
    results = []
    for k_s, (s, t, u) in enumerate(zip(sl, t_col, used)):
        r = SliceResult(interval=t, used=u)
        # петрофизика (считается для всех срезов)
        r.kvo = prm["swl"] if rg.kvo_a == 0 else -rg.kvo_a * s.phi + rg.kvo_b
        try:
            r.kno = prm["sorw_ow"] if rg.kno_a == 0 else (1 - r.kvo) * (1 - (rg.kno_a * math.log(s.perm) + rg.kno_b))
        except ValueError:
            r.kno = 0.0
        r.kvcr = prm["swcr_ow"] if rg.kvcr_a == 0 else -rg.kvcr_a * s.phi * 100 + rg.kvcr_b
        if s.is_gas:
            r.sg = s.so
        if not u:
            results.append(r)
            continue
        r.h_eff = _interval_value(case.intervals, t, "h_eff", d.h_eff)
        r.p_res = _interval_value(case.intervals, t, "p_res", d.p_res)
        if p_res_override is not None and p_res_override[k_s] > 0:
            r.p_res = p_res_override[k_s]
        flood = _interval_value(case.intervals, t, "flood_share", d.flood_share)
        r_e = _interval_value(case.intervals, t, "r_contour", d.r_contour)
        r.r_contour, r.flood_share = r_e, flood
        inj = _interval_value(case.intervals, t, "inj_ratio", d.inj_ratio)

        p = obj.pvt.calc(r.p_res, t_C)
        mu_o, mu_g, bo, bg, rs = p["mu_oil"], p["mu_gas"], p["bo"], p["bg"], p["rs"]
        mu_w = ph.mu_water_cP(r.p_res, t_C, obj.pvt_cfg["gamma_wat"])
        bw = ph.bw_m3m3(r.p_res, t_C)
        r.pvt = dict(mu_o=mu_o, mu_w=mu_w, mu_g=mu_g, bo=bo, bw=bw, bg=bg, rs=rs)

        sw = s.sw                                                   # AW
        if not s.is_gas:
            r.sg = 1 - obj.so(r.p_res) / (1 - prm["swi"]) * (1 - sw) - sw   # BI
        r.so = 1 - r.sg - sw                                        # BJ
        try:
            r.kro = ph.kro_baker(sw, r.kvo, r.sg, prm["sorg_og"], r.kno, prm["end_oil"], prm["exo_ow"], prm["exo_og"])
        except (ZeroDivisionError, ValueError):
            r.kro = 0.0
        if isinstance(r.kro, complex):
            r.kro = 0.0
        try:
            r.krg = ph.krg_baker(sw, r.kvo, r.so, r.sg, prm["sorg_og"], prm["end_gas"], prm["sgcr_og"], prm["exg_og"],
                                 r.kno, prm["swcr_ow"])
        except (ZeroDivisionError, ValueError):
            r.krg = 0.0
        if isinstance(r.krg, complex):
            r.krg = 0.0
        krw = ph.krw_baker(sw, r.kvcr, r.so, r.sg, prm["sorg_og"], prm["end_water"], r.kno, prm["exw_ow"], r.kvo)
        r.krw = 0.0 if (isinstance(krw, complex) or krw < 0) else krw
        r.fw = ph.fw_from_kr(mu_w, mu_o, bo, bw, r.krw, r.kro)      # AF
        r.mu = ph.mu_from_kr(mu_w, mu_o, r.krw, r.kro)              # BN
        r.ct = obj.ct(r.p_res)                                      # BO
        r.b = ph.b_liq(bo, bw, r.fw)                                # BP
        r.l_total = (len_in if t > 0 else len_out) if l_total is None else l_total   # BS
        r.l_section = s.h                                           # BT
        if case.well_type == "vertical":
            # расширение (в книге нет): пласт-срез вскрыт вертикальным стволом,
            # мощность — по вертикали (TVD), если задана; продуктивности слоёв складываются
            h_v = (s.bottom_tvd - s.top_tvd) if s.bottom_tvd > s.top_tvd else s.h
            r.k_section = 1.0
            r.pi = ph.get_j_obsolete(d.kv_kh, s.phi, 0.0, d.t_pss, r_e, d.r_w, h_v, s.perm, 0.0,
                                     r.mu, r.ct, r.b, 0.3)
        else:
            r.k_section = pallets.coefficient(r.l_total, r.l_section, r.h_eff)   # CA
            r.pi = ph.get_j_obsolete(d.kv_kh, s.phi, s.h, d.t_pss, r_e, d.r_w, r.h_eff, s.perm, 0.0,
                                     r.mu, r.ct, r.b, 0.3) * r.k_section   # BQ
        r.dp = ph.get_dp(prm["pb"], d.p_wf, r.fw, r.p_res)          # BR
        r.q_liq = r.pi * r.dp * case.k_success                      # AE
        r.q_oil = r.q_liq * rho * (1 - r.fw / 100)                  # AG
        r.gor = 0.0 if r.q_oil == 0 else ph.gor(rs, r.krg, r.kro, mu_o, bo, mu_g, bg)   # AH
        r.jd = ph.get_jd(r.mu, r.b, r.h_eff * s.perm, r.pi)         # AQ

        # запасы (AR:BA)
        bu = by_interval.get(t, 0.0)
        share = r.l_section / bu if bu else 0.0
        r_eq = math.sqrt((r_e * 2 * s.h + math.pi * r_e ** 2) / math.pi)
        n_dep = ph.get_n(s.phit, bo, rho, r_eq, sw, r.h_eff * (1 - flood)) * share   # AR
        n_disp = ph.get_n(s.phit, bo, rho, r_eq, sw, r.h_eff * flood) * share        # AS
        n_inj = inj * ph.get_n(s.phi, bo, rho, r_e, sw, r.h_eff * flood) * share     # AT
        k_disp = (1 - sw - prm["sorw_ow"]) / ((1 - sw - prm["sorw_ow"]) + prm["sorw_ow"])
        npz_disp, npz_inj = k_disp * n_disp, k_disp * n_inj                          # AU, AV
        npz_share = obj.npz_pallet(sw, 98.0)                                         # AY
        sg_p = obj.sg(r.p_res)
        niz_dep = n_dep * ((1 - sw - (sg_p / (1 - prm["swi"])) * (1 - sw))
                           - (1 - sw - (sg_pwf / (1 - prm["swi"])) * (1 - sw))) / (bo_pwf / obj.bo_curve_at(r.p_res))  # AZ
        niz_inj = (npz_disp + npz_inj) * npz_share                                   # BA
        r.oip = (n_dep + n_disp + n_inj) * rho                                       # AI
        r.rec = (niz_dep + niz_inj) * rho                                            # AJ
        results.append(r)
    return results


# --- Чтение листа «РИГИС» книги RE Tools ------------------------------------

def _f(v, default=0.0) -> float:
    return float(v) if isinstance(v, (int, float)) else default


def read_case_from_book(book, sheet: str = "РИГИС", wb=None) -> RigisCase:
    """Срезы (строки 73+) и настройки листа РИГИС книги RE Tools."""
    import warnings
    import openpyxl
    if wb is None:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = openpyxl.load_workbook(book, read_only=True, data_only=True)
    ws = wb[sheet]
    top = [list(r) for r in ws.iter_rows(min_row=1, max_row=18, max_col=43, values_only=True)]
    cell = lambda r, c: top[r - 1][c - 1]   # noqa: E731
    intervals = []
    for r in range(4, 14):
        n = cell(r, 2)
        if isinstance(n, (int, float)):
            intervals.append(IntervalParams(int(n), _f(cell(r, 3)), _f(cell(r, 4)), _f(cell(r, 7)), _f(cell(r, 8)),
                                            _f(cell(r, 9)), _f(cell(r, 10)), _f(cell(r, 11))))
    defaults = Defaults(*[_f(cell(15, c)) for c in range(3, 12)])
    regressions = Regressions(*[_f(cell(5, c)) for c in range(37, 44)])
    rules = CompletionRules(_f(cell(4, 19)), _f(cell(4, 20)), _f(cell(5, 20)), _f(cell(7, 20)), _f(cell(9, 20)),
                            _f(cell(11, 20)), _f(cell(12, 20)), _f(cell(13, 20)), _f(cell(14, 20)))
    slices = []
    for row in ws.iter_rows(min_row=73, max_col=17, values_only=True):
        if row[3] is None or not isinstance(row[3], (int, float)):
            if row[1] is None:
                break
            continue
        slices.append(Slice(str(row[1] or ""), str(row[2] or ""), _f(row[3]), _f(row[4]), _f(row[5]), _f(row[6]),
                            _f(row[7]), _f(row[8]), _f(row[9]), _f(row[10]), _f(row[11]), _f(row[12]), _f(row[13]),
                            _f(row[14]), str(row[15] or ""), str(row[16] or "")))
    return RigisCase(slices=slices, intervals=intervals, defaults=defaults, regressions=regressions, rules=rules,
                     k_success=_f(cell(18, 32), 1.0), object_name=str(cell(17, 28) or "PVT_MBE_Waterflood"))
