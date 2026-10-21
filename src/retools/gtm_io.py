# -*- coding: utf-8 -*-
"""Ввод-вывод рейтинга ГТМ и результата расчёта добычи (RE Tools).

Контракт движка расчёта добычи:
- лист «Рейтинг ГТМ» — входная таблица «скважина × параметры» (шапка в строке 6,
  данные с 7-й): местоположение, объект, ГТМ и вариант, дата запуска,
  заканчивание, фильтрационно-ёмкостные свойства, геометрия, режимы ППД,
  ограничения и уставки расчёта;
- лист «Результат расчёта» — месячные ряды (сводная таблица книги): Qн, Qж,
  обводнённость, приёмистость и накопленные отборы по датам.

Используется как источник данных для портируемого движка `production`
(спецификация шага по месяцам — docs/production_spec.md).
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import openpyxl

# Имена колонок рейтинга ГТМ (строка 6 книги) -> поля записи
FIELDS = {
    "A": "n", "B": "region", "C": "field", "D": "pad", "E": "well",
    "F": "target_zone", "G": "pvt_sheet", "H": "gtm", "I": "case",
    "J": "start_date", "K": "completion", "L": "purpose",
    "M": "h_eff", "N": "phi", "O": "perm", "P": "kv_kh", "Q": "sw_init",
    "R": "p_res_init", "S": "flood_share", "T": "l_hor", "U": "l_hor_eff",
    "Y": "frac_flag", "AF": "r_contour", "AG": "r_well", "AH": "skin",
    "AI": "p_wf_target", "AJ": "t_pss", "AK": "workover_time",
    "AL": "inj_start", "AM": "inj_prod_ratio", "AN": "inj_buf_p",
    "AQ": "inj_max_rate", "AR": "vrr_target", "AS": "aquifer_ratio",
    "AT": "q_min", "AU": "wc_limit", "AV": "qliq_dp_control",
    "AW": "jd_success", "AX": "operation_factor", "AY": "duration",
    "AZ": "oip", "BA": "pot_recovery", "BB": "rec_reserves",
}


@dataclass
class GtmWell:
    """Строка рейтинга ГТМ — скважина с параметрами для расчёта добычи."""
    n: int = 0
    region: str = ""
    field: str = ""
    pad: str = ""
    well: str = ""
    target_zone: str = ""
    pvt_sheet: str = ""
    gtm: str = ""
    case: str = ""
    start_date: object = None
    completion: str = ""
    purpose: str = ""
    h_eff: float = 0.0
    phi: float = 0.0
    perm: float = 0.0
    kv_kh: float = 0.0
    sw_init: float = 0.0
    p_res_init: float = 0.0
    flood_share: float = 0.0
    l_hor: float = 0.0
    l_hor_eff: float = 0.0
    frac_flag: float = 0.0
    r_contour: float = 0.0
    r_well: float = 0.0
    skin: float = 0.0
    p_wf_target: float = 0.0
    t_pss: float = 0.0
    workover_time: float = 0.0
    inj_start: object = None
    inj_prod_ratio: float = 0.0
    inj_buf_p: float = 0.0
    inj_max_rate: float = 0.0
    vrr_target: float = 0.0
    aquifer_ratio: float = 0.0
    q_min: float = 0.0
    wc_limit: float = 0.0
    qliq_dp_control: object = None
    jd_success: float = 0.0
    operation_factor: float = 0.0
    duration: float = 0.0
    oip: float = 0.0
    pot_recovery: float = 0.0
    rec_reserves: float = 0.0


def read_gtm(path: str | Path, sheet: str = "Рейтинг ГТМ") -> list[GtmWell]:
    """Прочитать рейтинг ГТМ (строки с непустым №)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(min_row=7, values_only=False):
        cellmap = {c.column_letter: c.value for c in row
                   if getattr(c, "column_letter", None) is not None}
        if cellmap.get("A") is None:
            if out:
                break                     # таблица закончилась
            continue
        kwargs = {}
        for col, name in FIELDS.items():
            v = cellmap.get(col)
            if v is not None:
                kwargs[name] = v
        out.append(GtmWell(**kwargs))
    wb.close()
    return out


def read_result(path: str | Path, sheet: str = "Результат расчёта") -> list[dict]:
    """Прочитать месячные ряды результата: даты + суммы/средние книги.

    Первая колонка — дата, дальше пары «метрика -> значение» по заголовкам
    строки 1 (Сумма по полю Qн, Среднее по полю Обв-ть, ...).
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(min_row=1, values_only=True))
    headers = [str(v) if v is not None else "" for v in rows[0]]
    out = []
    for r in rows[1:]:
        if r[0] is None:
            break
        rec = {"date": r[0]}
        for h, v in zip(headers[1:], r[1:]):
            if h and v is not None:
                rec[h] = v
        out.append(rec)
    wb.close()
    return out


# --- входы рейтинга и лист объекта для порта «Расчёта добычи» (gtm_calc) ---------------------

RATING_INPUT_COLS = ("A B C D E F G H I J K L M N O P Q R S T U V W X Y Z AA AB AC AD AE AF AG AH AI AJ AK AL AM "
                     "AN AO AP AQ AR AS AT AU AV AW AX AY").split()


def _load_values(path):
    import warnings
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return openpyxl.load_workbook(path, read_only=True, data_only=True)


def read_rating_rows(path=None, sheet: str = "Рейтинг ГТМ", wb=None) -> list[dict]:
    """Входы рейтинга ГТМ (A:AY, строки 7+, пока заполнен №) — словари по буквам столбцов."""
    from openpyxl.utils import get_column_letter
    wb = wb or _load_values(path)
    ws = wb[sheet]
    out = []
    for row in ws.iter_rows(min_row=7, max_col=len(RATING_INPUT_COLS), values_only=True):
        if row[0] is None:
            if out:
                break
            continue
        out.append({get_column_letter(i + 1): v for i, v in enumerate(row) if v is not None})
    return out


def read_object_sheet(path=None, sheet: str = "PVT_MBE_Waterflood", wb=None):
    """Лист объекта книги как gtm_calc.ObjectSheet (параметры, строки 3–38, палетки)."""
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as gl
    from .gtm_calc import ObjectSheet
    wb = wb or _load_values(path)
    ws = wb[sheet]
    grid = {}
    for row in ws.iter_rows(min_row=1, max_row=185, min_col=1, max_col=64):
        for c in row:
            if c.value is not None:
                grid[(c.row, c.column)] = c.value

    def block(c0, c1, r_axis, r0, r1, ycol):
        return {"x": [grid.get((r_axis, c)) for c in range(ci(c0), ci(c1) + 1)],
                "y": [grid.get((r, ci(ycol))) for r in range(r0, r1 + 1)],
                "f": [[grid.get((r, c)) for c in range(ci(c0), ci(c1) + 1)] for r in range(r0, r1 + 1)]}

    book = {"D": {str(r): grid.get((r, 4)) for r in range(2, 39)},
            "aux": {k: grid.get((int(k[2:]), ci(k[:2]))) for k in ("AZ38", "BA38", "BA21", "BA27")},
            "table": {gl(c): [grid.get((r, c)) for r in range(3, 39)] for c in range(ci("E"), ci("U") + 1)},
            "pallets": {"mu": block("D", "O", 69, 70, 105, "C"), "kro_mix": block("R", "AC", 69, 70, 105, "Q"),
                        "krw_mix": block("AF", "AQ", 69, 70, 105, "AE"), "sg": block("D", "O", 108, 109, 144, "C"),
                        "wc_mix": block("R", "AC", 108, 109, 144, "Q"), "npz": block("AF", "AQ", 108, 109, 144, "AE"),
                        "wc": block("D", "O", 147, 148, 183, "C")}}
    obj = ObjectSheet.from_book(book)
    obj.name = sheet
    return obj
