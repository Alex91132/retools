# -*- coding: utf-8 -*-
"""Выгрузка подбора компоновки в Excel: компоновка, изоляция, срезы, кривая."""
from __future__ import annotations

import io
from dataclasses import asdict

from .completion import Completion, liner


def completion_xlsx(case, comp: Completion, curve: list | None = None, title: str = "") -> bytes:
    """Выгрузка компоновки одного ствола ГС."""
    from .completion import WellCompletion
    wc = WellCompletion([comp], [case.slices[0].well if case.slices else ""], comp.q_liq, comp.q_oil,
                        comp.water_cut, comp.l_total, comp.max_wc, comp.iterations)
    return well_xlsx([case], wc, curve, title)


def well_xlsx(cases, wc, curve: list | None = None, title: str = "", forecasts: dict | None = None,
              economics: list | None = None) -> bytes:
    """Выгрузка компоновки скважины (ГС или МЗС): компоновка, изоляция, срезы, кривая, прогноз,
    экономика (строки-словари по вариантам), настройки."""
    import openpyxl
    from openpyxl.styles import Font

    wb = openpyxl.Workbook()
    bold = Font(bold=True)
    multi = len(cases) > 1

    def header(ws, cols):
        ws.append(cols)
        for c in ws[ws.max_row]:
            c.font = bold

    ws = wb.active
    ws.title = "Компоновка"
    ws.append([title or ("Подбор компоновки МЗС по РИГИС" if multi else "Подбор компоновки ГС по РИГИС")])
    ws["A1"].font = bold
    ws.append(["Предел обводнённости, %", wc.max_wc])
    ws.append(["Qн, т/сут", round(wc.q_oil, 2)])
    ws.append(["Qж, м³/сут", round(wc.q_liq, 2)])
    ws.append(["Обводнённость, %", round(wc.water_cut, 1)])
    ws.append(["Длина фильтров (работающих срезов), м", round(wc.l_total, 1)])
    if multi:
        ws.append([])
        header(ws, ["Ствол", "Qн, т/сут", "Qж, м³/сут", "Обв., %", "Фильтров"])
        for name, c in zip(wc.names, wc.laterals):
            ws.append([name, round(c.q_oil, 2), round(c.q_liq, 2), round(c.water_cut, 1), len(c.filters)])
    ws.append([])
    header(ws, ["Ствол", "Элемент", "Кровля, м", "Подошва, м", "Длина, м"])
    for name, case, c in zip(wc.names, cases, wc.laterals):
        for e in liner(c, case.rules):
            ws.append([name, e.kind, round(e.top, 2), round(e.bottom, 2), round(e.length, 2)])
    ws.append([])
    header(ws, ["Ствол", "Фильтр", "Кровля, м", "Подошва, м", "Длина, м", "Qж, м³/сут", "Qн, т/сут", "Обв., %"])
    for name, c in zip(wc.names, wc.laterals):
        for i, f in enumerate(c.filters, 1):
            ws.append([name, i, round(f.top, 2), round(f.bottom, 2), round(f.length, 2), round(f.q_liq, 2),
                       round(f.q_oil, 2), round(f.water_cut, 1)])

    ws = wb.create_sheet("Изоляция")
    header(ws, ["Ствол", "Кровля, м", "Подошва, м", "Длина, м", "Причина", "Потенциал Qж, м³/сут",
                "Потенциал Qн, т/сут", "Обв., %", "Пакеры (кровля–подошва)"])
    for name, c in zip(wc.names, wc.laterals):
        for iso in c.isolations:
            ws.append([name, round(iso.top, 2), round(iso.bottom, 2), round(iso.bottom - iso.top, 2), iso.reason,
                       round(iso.q_liq, 2), round(iso.q_oil, 2), round(iso.water_cut, 1),
                       "; ".join(f"{a:.1f}–{b:.1f}" for a, b in iso.packers)])

    ws = wb.create_sheet("Срезы")
    header(ws, ["Ствол", "Кровля", "Подошва", "Мощность", "УЭС", "Кп", "Кпр", "Кв", "Насыщение", "Коллектор",
                "В работе", "Интервал", "Qж", "Fw, %", "Qн", "ГФ", "ПИ", "dP", "К секции", "НГЗ, тыс.т", "НИЗ, тыс.т"])
    for name, case, comp in zip(wc.names, cases, wc.laterals):
        inside = lambda s: any(f.top <= s.top and s.bottom <= f.bottom for f in comp.filters)   # noqa: E731
        for s, r in zip(case.slices, comp.slices):
            ws.append([name, s.top, s.bottom, s.h, s.res, s.phi, s.perm, s.sw, s.saturation, s.collector,
                       "фильтр" if inside(s) else "", r.interval, r.q_liq, r.fw, r.q_oil, r.gor, r.pi, r.dp,
                       r.k_section, r.oip, r.rec])

    if curve:
        ws = wb.create_sheet("Кривая")
        header(ws, ["Предел обв., %", "Qн, т/сут", "Qж, м³/сут", "Обв., %", "Фильтров", "Длина, м"])
        for c in curve:
            n_f = len(c.filters) if hasattr(c, "filters") else sum(len(x.filters) for x in c.laterals)
            ws.append([c.max_wc, round(c.q_oil, 2), round(c.q_liq, 2), round(c.water_cut, 1), n_f, round(c.l_total, 1)])

    if forecasts:
        ws = wb.create_sheet("Прогноз")
        header(ws, ["Вариант", "Месяц", "Рпл, атм", "Qн, т/сут", "Qж, м³/сут", "Обв., %", "Накоплено нефти, тыс. т",
                    "Накоплено жидкости, тыс. м³", "Sw средн.", "Переток Qж, м³/сут", "Доля перетока", "В работе"])
        for name, f in forecasts.items():
            for m in f.months:
                ws.append([name, m.month, round(m.p_res, 2), round(m.q_oil, 2), round(m.q_liq, 2), round(m.water_cut, 1),
                           round(m.cum_oil, 3), round(m.cum_liq, 3), round(m.sw, 4), round(m.q_liq_leak, 2),
                           round(m.leak, 4), "да" if m.working else "нет"])

    if economics:
        ws = wb.create_sheet("Экономика")
        header(ws, list(economics[0]))
        for r in economics:
            ws.append(list(r.values()))

    ws = wb.create_sheet("Настройки")
    case = cases[0]
    for name, obj in (("Правила компоновки", case.rules), ("Параметры по умолчанию", case.defaults),
                      ("Регрессии петрофизики", case.regressions)):
        ws.append([name])
        ws[ws.max_row][0].font = bold
        for k, v in asdict(obj).items():
            ws.append([k, v])
        ws.append([])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
