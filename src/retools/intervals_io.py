# -*- coding: utf-8 -*-
"""Чтение и агрегация листа «РИГИС» расчётника RE Tools.

Лист устроен так (проверено по книге RE Tools 18.08.23):
- строки 4-13: интервалы перфорации «Выбор» (№, кровля, подошва, эфф. мощность,
  нач. Рпл, доля заводняемых толщин, R контура, соотношение ППД/доб);
- строка 14-15: итоги по скважине и параметры по умолчанию;
- строки 73-451: петрофизика по глубинным срезам: колонка D — глубина,
  Q — флаг коллектора («К»), R — толщина среза, отнесённая к интервалу,
  AK — код объекта (в формулах книги — именованный диапазон «Компкол», ключ
  агрегации SUMIF для НГЗ), AD..AM — расчётные свойства среза
  (УЭС, Qж, Fw, Qн, ГФ, НГЗ, извлекаемые, эфф. мощность, Рпл, ...).

Модуль читает книгу и воспроизводит агрегацию интервалов (суммы НГЗ,
эффективной мощности, средневзвешенные параметры) без Excel.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# Колонки листа «РИГИС» (1-индексные)
COL = dict(depth=4, collector=17, thick=18, code=37,      # D, Q, R, AK
           ues=30, qliq=31, fw=32, qoil=33, gor=34,        # AD..AH
           oip=35, rec_oip=36, h_eff=38, p_res=39)         # AI, AJ, AL, AM
SLICE_ROWS = (73, 451)


@dataclass
class IntervalChoice:
    """Интервал перфорации из блока «Выбор» (строки 4-13 листа)."""
    n: int
    top: float
    bottom: float
    h_eff: float = 0.0
    p_res: float = 0.0
    flood_share: float = 0.0
    r_contour: float = 0.0
    inj_prod_ratio: float = 0.0


@dataclass
class RigisData:
    """Прочитанный лист РИГИС: интервалы, срезы, параметры по умолчанию."""
    intervals: list[IntervalChoice] = field(default_factory=list)
    slices: list[dict] = field(default_factory=list)
    defaults: dict = field(default_factory=dict)

    def slices_of(self, interval_code: int) -> list[dict]:
        """Срезы, относящиеся к интервалу (по коду объекта «Компкол»)."""
        return [s for s in self.slices if s["code"] == interval_code]

    def aggregate(self, interval_code: int) -> dict:
        """Агрегация интервала: суммы и средневзвешенные по срезам.

        Воспроизводит итоговые формулы книги: НГЗ интервала — SUMIF по коду
        объекта (колонка AI), эфф. мощность — сумма толщин срезов (AL),
        нач. Рпл — средневзвешенное по мощности.
        """
        ss = self.slices_of(interval_code)
        if not ss:
            return {"h_eff": 0.0, "oip": 0.0, "p_res": 0.0,
                    "qliq": 0.0, "qoil": 0.0, "n_slices": 0}
        h = sum(s["h_eff"] for s in ss)
        if h > 0:
            p_res = sum(s["p_res"] * s["h_eff"] for s in ss) / h
        else:
            p_res = sum(s["p_res"] for s in ss) / len(ss)
        return {
            "h_eff": h,
            "oip": sum(s["oip"] for s in ss),
            "p_res": p_res,
            "qliq": sum(s["qliq"] for s in ss),
            "qoil": sum(s["qoil"] for s in ss),
            "n_slices": len(ss),
        }

    def totals(self) -> dict:
        """Итог по скважине: суммы по активным интервалам (строка 14 книги)."""
        per = [self.aggregate(i.n) for i in self.intervals if i.h_eff > 0]
        return {
            "oip": sum(p["oip"] for p in per),
            "h_eff": sum(p["h_eff"] for p in per),
            "qliq": sum(p["qliq"] for p in per),
            "n_intervals": len(per),
        }


def read_rigis(path: str | Path, sheet: str = "РИГИС") -> RigisData:
    """Прочитать лист РИГИС из книги RE Tools (значения, без формул)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]

    data = RigisData()
    for ridx, row in enumerate(ws.iter_rows(min_row=4, max_row=15, values_only=True), start=4):
        if ridx <= 13 and isinstance(row[1], (int, float)):
            data.intervals.append(IntervalChoice(
                n=int(row[1]), top=row[2] or 0.0, bottom=row[3] or 0.0,
                h_eff=row[6] or 0.0, p_res=row[7] or 0.0,
                flood_share=row[8] or 0.0, r_contour=row[9] or 0.0,
                inj_prod_ratio=row[10] or 0.0))
        if ridx == 15:  # параметры по умолчанию
            data.defaults = {
                "kv_kh": row[2], "r_well": row[3], "t_pss": row[4],
                "p_wf_target": row[5], "h_eff": row[6], "p_res": row[7],
                "flood_share": row[8], "r_contour": row[9],
                "inj_prod_ratio": row[10]}
    # петрофизика: срезы
    for row in ws.iter_rows(min_row=SLICE_ROWS[0], max_row=SLICE_ROWS[1],
                            values_only=True):
        depth = row[COL["depth"] - 1]
        if not isinstance(depth, (int, float)):
            continue
        data.slices.append({
            "depth": float(depth),
            "collector": str(row[COL["collector"] - 1] or ""),
            "thick_to_interval": float(row[COL["thick"] - 1] or 0),
            # код объекта может быть числом (№ интервала) или текстом (имя листа
            # PVT); числовые сравниваются с № интервала, текстовые не матчатся
            "code": row[COL["code"] - 1] if row[COL["code"] - 1] is not None else -1,
            "ues": float(row[COL["ues"] - 1] or 0),
            "qliq": float(row[COL["qliq"] - 1] or 0),
            "fw": float(row[COL["fw"] - 1] or 0),
            "qoil": float(row[COL["qoil"] - 1] or 0),
            "gor": float(row[COL["gor"] - 1] or 0),
            "oip": float(row[COL["oip"] - 1] or 0),
            "rec_oip": float(row[COL["rec_oip"] - 1] or 0),
            "h_eff": float(row[COL["h_eff"] - 1] or 0),
            "p_res": float(row[COL["p_res"] - 1] or 0),
        })
    wb.close()
    return data
