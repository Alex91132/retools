# -*- coding: utf-8 -*-
"""Страница Streamlit «Подбор компоновки по РИГИС».

Сценарий: загрузить РИГИС (выгрузка Excel или книга RE Tools) → задать зоны
пласта и правила компоновки → «Подобрать» → планшет с фильтрами и пакерами,
таблицы фильтров и изоляции, кривая «Qн от предела обводнённости» → выгрузка
компоновки в Excel и проекта в JSON.
"""
from __future__ import annotations

import json
import tempfile
from dataclasses import asdict, replace
from pathlib import Path

import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from plotly.subplots import make_subplots

from .completion import evaluate_multi, liner, manual_plan, split_laterals, wc_curve_multi
from .economics import Economics, capex, cash_flow, npv
from .forecast import forecast
from .lifetime import OBJECTIVES, plan_lifetime
from .field_object import read_field_object, read_jd_pallets
from .object_builder import build_jd_pallets, build_object, jd_settings_from_object, mbe_table
from .report import well_xlsx
from .rigis import CompletionRules, Defaults, IntervalParams, Regressions, RigisCase, compute, read_case_from_book
from .rigis_io import case_from_dict, case_to_dict, load_project, read_rigis_table, save_project

DEMO = Path(__file__).resolve().parents[2] / "tests" / "fixtures" / "rigis_4210H_project.json"
SAT_COLORS = {"нефть": "#2e9d4f", "нефть+вода": "#7fb46a", "вода": "#2f7fd8", "газ": "#d64545",
              "разгазирование": "#e39b32"}
LINER_COLORS = {"фильтр": "#2e9d4f", "эластомер": "#e07b00", "глухая труба": "#9aa0a6"}


def _tmp(upload) -> Path:
    suffix = Path(upload.name).suffix
    f = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    f.write(upload.getbuffer())
    f.close()
    return Path(f.name)


def _load_book(path: Path):
    """Книга RE Tools: срезы, настройки, объект и палетки (макросы не выполняются)."""
    import warnings
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    case = read_case_from_book(path, wb=wb)
    obj = read_field_object(path, case.object_name, wb=wb)
    pal = read_jd_pallets(path, wb=wb)
    return case, obj, pal


def _sources():
    st.subheader("1. Данные")
    src = st.radio("Источник", ["Выгрузка РИГИС (Excel)", "Книга RE Tools (.xlsm)", "Файл проекта (.json)",
                                "Пример: скважина 4210H", "Пример МЗС: 2 ствола (синтетика)"], horizontal=True, key="src")
    state = st.session_state
    if src == "Пример: скважина 4210H":
        if st.button("Открыть пример", key="open_demo") or "case" not in state:
            if DEMO.exists():
                state.case, state.obj, state.pal = load_project(DEMO)
                state.source = "пример 4210H"
            else:
                st.warning("Файл примера не найден")
    elif src == "Пример МЗС: 2 ствола (синтетика)":
        if st.button("Открыть пример МЗС", key="open_mzs") or state.get("source") != "пример МЗС":
            if DEMO.exists():
                case, state.obj, state.pal = load_project(DEMO)
                # второй ствол — синтетика: те же срезы, водонасыщенность выше на 0,08
                lat2 = [replace(sl, well="4210H/2", sw=min(sl.sw + 0.08, 0.95), so=max(sl.so - 0.08, 0.0))
                        for sl in case.slices]
                state.case = replace(case, slices=list(case.slices) + lat2)
                state.source = "пример МЗС"
    elif src == "Книга RE Tools (.xlsm)":
        up = st.file_uploader("Книга RE Tools", type=["xlsm", "xlsx"], key="up_book")
        if up is not None and state.get("source") != up.name:
            with st.spinner("Читаю книгу (макросы не выполняются)…"):
                state.case, state.obj, state.pal = _load_book(_tmp(up))
            state.source = up.name
    elif src == "Файл проекта (.json)":
        up = st.file_uploader("Проект", type=["json"], key="up_proj")
        if up is not None and state.get("source") != up.name:
            state.case, state.obj, state.pal = load_project(_tmp(up))
            state.source = up.name
    else:
        c1, c2 = st.columns([2, 1])
        with c1:
            up = st.file_uploader("Выгрузка РИГИС по скважине", type=["xlsx", "xlsm"], key="up_rigis")
        with c2:
            obj_src = st.file_uploader("Объект разработки и палетки: проект .json или книга RE Tools", type=["json", "xlsm"],
                                       key="up_obj")
            st.caption("Без файла — объект и палетки примера 4210H.")
        if up is not None and state.get("source") != up.name:
            path = _tmp(up)
            sheet = None
            try:
                import openpyxl
                names = openpyxl.load_workbook(path, read_only=True).sheetnames
                if len(names) > 1:
                    sheet = st.selectbox("Лист", names, key="rigis_sheet")
                slices, mapping = read_rigis_table(path, sheet=sheet)
            except ValueError as e:
                st.error(str(e))
                return
            if obj_src is not None:
                p = _tmp(obj_src)
                base, obj, pal = load_project(p) if p.suffix == ".json" else _load_book(p)
            else:
                base, obj, pal = load_project(DEMO)
            state.case = replace(base, slices=slices, intervals=[IntervalParams(i) for i in range(1, 11)])
            state.obj, state.pal = obj, pal
            state.source = up.name
            state.mapping = mapping
    if "case" in state:
        c = state.case
        st.caption(f"Источник: **{state.get('source', '')}** · срезов {len(c.slices)} · "
                   f"{c.slices[0].top:.1f}–{c.slices[-1].bottom:.1f} м · объект «{state.obj.name}»")
        if state.get("mapping"):
            with st.expander("Распознанные столбцы выгрузки"):
                st.table(pd.DataFrame([{"поле": k, "столбец": v} for k, v in state.mapping.items()]))


OBJ_FIELDS = [
    ("Пласт", [("p_init", "Рпл нач., атм", 1.0), ("depth", "Глубина пласта, м", 10.0), ("t_C", "Т пласта, °C", 1.0),
               ("c_f", "Сжимаемость породы, 1/атм", 1e-5), ("gas_cap_m", "m газовой шапки", 0.05),
               ("swl", "Swl", 0.01), ("swi", "Swi (среднее)", 0.01)]),
    ("Флюид (PVT)", [("gamma_gas", "γ газа", 0.01), ("rho_oil", "γ нефти", 0.01), ("gamma_wat", "γ воды", 0.01),
                     ("rsi", "Rsb, м³/м³", 1.0), ("pb", "Pb, атм", 1.0), ("bob", "Bob, м³/м³", 0.01),
                     ("muob", "μob, сПз", 0.1), ("corr_set", "Корреляции (0 Standing, 1 McCain)", 1.0)]),
    ("ОФП (Бейкер)", [("sorw_ow", "Sorw", 0.01), ("sorg_og", "Sorg", 0.01), ("swcr_ow", "Swcr", 0.01),
                      ("sgcr_og", "Sgcr", 0.01), ("end_oil", "Kro концевая", 0.01), ("end_water", "Krw концевая", 0.01),
                      ("end_gas", "Krg концевая", 0.01), ("exo_ow", "Exo_ow", 0.1), ("exw_ow", "Exw_ow", 0.1),
                      ("exo_og", "Exo_og", 0.1), ("exg_og", "Exg_og", 0.1)]),
    ("Единицы потока (палетки секций и заводнения)", [("heff_fu1", "h ЕП-1, м", 1.0), ("k_fu1", "k ЕП-1, мД", 1.0),
                                             ("phi_fu1", "φ ЕП-1", 0.01), ("heff_fu2", "h ЕП-2, м", 1.0),
                                             ("k_fu2", "k ЕП-2, мД", 1.0), ("phi_fu2", "φ ЕП-2", 0.01)]),
]


def _object_editor():
    state = st.session_state
    obj = state.obj
    with st.expander(f"Объект разработки и палетки секций ГС — сейчас: «{obj.name}»", expanded=False):
        st.caption("По умолчанию объект и палетки берутся из книги или проекта. Для другого месторождения задайте "
                   "параметры и постройте объект: таблица матбаланса (Np, So, Sg, Ct от давления) решается заново, "
                   "палетки секций ГС генерируются по Get_J_obsolete, палетка «доля НПЗ при предельной обводнённости» — "
                   "по модели вытеснения двух единиц потока (порт макроса FDP_Functions).")
        prm = dict(obj.params)
        with st.form("obj_form"):
            new = {}
            for title, fields in OBJ_FIELDS:
                st.markdown(f"**{title}**")
                cols = st.columns(4)
                for i, (key, label, step) in enumerate(fields):
                    v = prm.get(key)
                    v = float(v) if isinstance(v, (int, float)) else 0.0
                    new[key] = cols[i % 4].number_input(label, value=v, step=step, format="%.6g", key=f"obj_{key}")
            st.markdown("**Генерация палеток секций**")
            c = st.columns(4)
            t_p = c[0].number_input("Время, сут", value=60.0, step=5.0)
            r_e = c[1].number_input("Rконт, м", value=350.0, step=10.0)
            r_w = c[2].number_input("Rскв, м", value=0.07, step=0.01, format="%.3f")
            kv_kh = c[3].number_input("Kv/Kh", value=0.1, step=0.01, format="%.3f")
            submitted = st.form_submit_button("Построить объект и палетки из параметров", type="primary")
        if submitted:
            prm.update(new)
            prm["corr_set"] = int(round(prm["corr_set"]))
            with st.spinner("Решаю матбаланс и строю палетки…"):
                built = build_object(prm, name="из параметров")
                settings = jd_settings_from_object(built, t_p=t_p, r_e=r_e, r_w=r_w, kv_kh=kv_kh)
                state.obj, state.pal = built, build_jd_pallets(settings)
            st.success(f"Объект построен; палетки: μ смеси {settings.mu:.2f} сПз, k {settings.k:.1f} мД, "
                       f"φ {settings.phi:.3f}, Ct {settings.c_t:.2e} 1/атм")
        rows = mbe_table(state.obj.params)
        fig = make_subplots(rows=1, cols=3, subplot_titles=("Отбор Np, % НГЗ", "Sg, So", "Ct, 1/атм"))
        ps = [r.p for r in rows]
        fig.add_trace(go.Scatter(x=ps, y=[r.np_pct for r in rows], name="Np", line=dict(color="#2e9d4f")), 1, 1)
        fig.add_trace(go.Scatter(x=ps, y=[r.sg for r in rows], name="Sg", line=dict(color="#d64545")), 1, 2)
        fig.add_trace(go.Scatter(x=ps, y=[r.so for r in rows], name="So", line=dict(color="#2e9d4f", dash="dot")), 1, 2)
        fig.add_trace(go.Scatter(x=ps, y=[r.ct for r in rows], name="Ct", line=dict(color="#5b6abf")), 1, 3)
        fig.update_yaxes(type="log", row=1, col=3)
        fig.update_xaxes(title_text="Рпл, атм")
        fig.update_layout(height=280, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.35))
        st.plotly_chart(fig, use_container_width=True)
        st.caption("Матбаланс текущего объекта, решённый заново. В книге ниже ~124 атм «Подбор параметра» не "
                   "сошёлся (невязка до −830), поэтому кривые книги там отличаются.")


def _settings(case: RigisCase) -> RigisCase:
    st.subheader("2. Настройки")
    c1, c2 = st.columns([3, 2])
    with c1:
        st.markdown("**Зоны пласта** (параметры по интервалам глубин; пусто — по умолчанию)")
        df = pd.DataFrame([asdict(i) for i in case.intervals]).rename(columns={
            "n": "№", "top": "Кровля", "bottom": "Подошва", "h_eff": "h эфф, м", "p_res": "Рпл нач, атм",
            "flood_share": "Доля заводн.", "r_contour": "Rконт, м", "inj_ratio": "ППД/доб"})
        ed = st.data_editor(df, hide_index=True, num_rows="fixed", key="zones", use_container_width=True)
        intervals = [IntervalParams(int(r["№"]), float(r["Кровля"] or 0), float(r["Подошва"] or 0),
                                    float(r["h эфф, м"] or 0), float(r["Рпл нач, атм"] or 0), float(r["Доля заводн."] or 0),
                                    float(r["Rконт, м"] or 0), float(r["ППД/доб"] or 0)) for _, r in ed.iterrows()]
        d = case.defaults
        st.markdown("**По умолчанию и скважина**")
        a, b, cc, dd, e = st.columns(5)
        p_wf = a.number_input("Рзаб целевое, атм", value=float(d.p_wf), step=1.0)
        h_eff = b.number_input("h эфф, м", value=float(d.h_eff), step=0.5)
        p_res = cc.number_input("Рпл, атм", value=float(d.p_res), step=1.0)
        r_c = dd.number_input("Rконт, м", value=float(d.r_contour), step=10.0)
        t_pss = e.number_input("Тзап, сут", value=float(d.t_pss), step=1.0)
        defaults = replace(d, p_wf=p_wf, h_eff=h_eff, p_res=p_res, r_contour=r_c, t_pss=t_pss)
    with c2:
        r = case.rules
        st.markdown("**Правила компоновки**")
        a, b = st.columns(2)
        max_wc = a.number_input("Предел обводнённости, %", 0.0, 100.0, float(r.max_wc), 1.0)
        max_filters = b.number_input("Макс. число фильтров", 1, 30, 10)
        min_filter = a.number_input("Мин. длина фильтра, м", 1.0, 500.0, float(r.min_filter), 1.0)
        min_gap = b.number_input("Мин. расстояние между фильтрами, м", 1.0, 500.0, float(r.min_gap), 1.0)
        elastomer = a.number_input("Длина эластомера, м", 0.5, 50.0, float(r.elastomer), 0.5)
        offset = b.number_input("Отступ эластомера от фильтра, м", 0.0, 50.0, float(r.elastomer_offset), 0.5)
        pgmc = a.number_input("Подошва ПГМЦ, м (выше — не работает)", value=float(r.pgmc_bottom), step=1.0)
        well_type = b.selectbox("Тип скважины", ["ГС / МЗС (секции ствола)", "Вертикальная / наклонная (перфорация)"],
                                index=0 if case.well_type == "horizontal" else 1)
        wc_mode = a.selectbox("Предел обводнённости", ["для каждого фильтра (как в книге)", "для скважины в целом"])
        gas_offset = b.number_input("Отступ фильтров от газа, м", 0.0, 500.0, 0.0, 5.0)
        filter_cost = a.number_input("Мин. прирост Qн на фильтр, т/сут", 0.0, 500.0, 0.0, 1.0,
                                     help="Фильтр ставится, только если добавляет больше; 0 — как в книге")
        st.session_state["filter_cost"] = filter_cost
        st.markdown("**Контакты** (абс. отметка, м; 0 — не учитывать)")
        g1, g2, g3, g4 = st.columns(4)
        goc = g1.number_input("ГНК, АО", 0.0, 10000.0, 0.0, 1.0)
        goc_off = g2.number_input("Отступ от ГНК, м", 0.0, 100.0, 0.0, 0.5)
        owc = g3.number_input("ВНК, АО", 0.0, 10000.0, 0.0, 1.0)
        owc_off = g4.number_input("Отступ от ВНК, м", 0.0, 100.0, 0.0, 0.5)
        st.session_state["contacts"] = dict(goc_tvd=goc or None, goc_offset=goc_off, owc_tvd=owc or None,
                                            owc_offset=owc_off)
        use_bar = st.checkbox("Барьер под стволом (прослой)", value=False,
                              help="Непроницаемый или плотный прослой (глина, плотняк) между стволом и ВНК/ГНК "
                                   "останавливает конус: вода снизу (газ сверху) проходит через него не больше "
                                   "Дарси-потока при депрессии среза, остальная вода обходит "
                                   "барьер и приходит сбоку. Отметки — абсолютные, как у ВНК.")
        if use_bar:
            b1, b2, b3 = st.columns(3)
            b_top = b1.number_input("Кровля барьера, АО", 0.0, 10000.0, float(owc - 1.5 if owc else (goc + 0.5 if goc else 0.0)),
                                    0.5, help="Барьер должен лежать целиком между стволом и ВНК (или между ГНК и стволом)")
            b_th = b2.number_input("Толщина, м", 0.1, 100.0, 1.0, 0.5)
            b_k = b3.number_input("Проницаемость, мД", 0.0, 1000.0, 0.0, 0.001, format="%.4f",
                                  help="0 — непроницаемый")
            st.session_state["barrier"] = dict(top_tvd=b_top, thickness=b_th, perm=b_k) if b_top > 0 else None
        else:
            st.session_state["barrier"] = None
        st.session_state["wc_mode"] = "filter" if wc_mode.startswith("для каждого") else "well"
        st.session_state["gas_offset"] = gas_offset
        x1, x2 = st.columns([3, 2])
        crossflow = x1.checkbox("Перетоки по пласту вдоль ствола", value=True,
                                help="Изолированный срез отдаёт к ближайшему фильтру долю потенциала через "
                                     "соседние срезы пласта; неколлектор — барьер. В книге этого нет.")
        trans_mult = x2.number_input("Проводимость вдоль ствола, доля", 0.0, 1.0, 1.0, 0.05,
                                     help="Калибровка: 1 — пласт однородный, меньше — есть непроницаемые прослои")
        st.session_state["crossflow"] = dict(crossflow=crossflow, trans_mult=trans_mult)
        use_regr = st.checkbox("Петрофизика по регрессиям объекта (Кво, Кно, Квкрит от Кп и Кпр)",
                               value=any(asdict(case.regressions).values()))
        rules = replace(r, max_wc=max_wc, min_filter=min_filter, min_gap=min_gap, elastomer=elastomer,
                        elastomer_offset=offset, pgmc_bottom=pgmc)
    regr = case.regressions
    if not use_regr:
        regr = Regressions()
    elif not any(asdict(regr).values()):
        regr = st.session_state.get("regr_saved", Regressions())
    if any(asdict(case.regressions).values()):
        st.session_state["regr_saved"] = case.regressions
    st.session_state["max_filters"] = int(max_filters)
    return replace(case, intervals=intervals, defaults=defaults, rules=rules, regressions=regr,
                   well_type="horizontal" if well_type.startswith("ГС") else "vertical")


def _planshet(case, comp, title):
    sl = case.slices
    res = comp.slices
    mid = [(s.top + s.bottom) / 2 for s in sl]
    h = [s.bottom - s.top for s in sl]
    fig = make_subplots(rows=1, cols=7, shared_yaxes=True, horizontal_spacing=0.015,
                        column_widths=[0.07, 0.15, 0.15, 0.13, 0.18, 0.14, 0.12],
                        subplot_titles=("Насыщ.", "УЭС, Омм", "Кп, Кв", "Кпр, мД", "Qн / Qж по срезам", "Fw, %",
                                        "Компоновка"))
    sat = [s.saturation.strip() for s in sl]
    colors = [SAT_COLORS.get(x, "#c9ccd1") if s.collector.strip() == "К" else "#e9eaec" for x, s in zip(sat, sl)]
    hover = [f"{s.top:.1f}–{s.bottom:.1f} м<br>{x or '—'}, {s.collector}<br>Кп {s.phi:.3f} Кв {s.sw:.3f} Кпр {s.perm:.0f}"
             for s, x in zip(sl, sat)]
    fig.add_trace(go.Bar(x=[1] * len(sl), y=mid, width=h, orientation="h", marker_color=colors, hovertext=hover,
                         hoverinfo="text", showlegend=False), 1, 1)
    step_y = [v for s in sl for v in (s.top, s.bottom)]
    step = lambda vals: [v for x in vals for v in (x, x)]   # noqa: E731
    fig.add_trace(go.Scatter(x=step([max(s.res, 0.1) for s in sl]), y=step_y, mode="lines", line=dict(color="#95a5a6", width=1),
                             name="УЭС", showlegend=False), 1, 2)
    fig.add_trace(go.Scatter(x=step([s.phi for s in sl]), y=step_y, mode="lines", line=dict(color="#8e44ad", width=1.2),
                             name="Кп"), 1, 3)
    fig.add_trace(go.Scatter(x=step([s.sw for s in sl]), y=step_y, mode="lines", line=dict(color="#2f7fd8", width=1.2),
                             name="Кв"), 1, 3)
    fig.add_trace(go.Scatter(x=step([max(s.perm, 0.01) for s in sl]), y=step_y, mode="lines",
                             line=dict(color="#b8860b", width=1), name="Кпр", showlegend=False), 1, 4)
    fig.add_trace(go.Bar(x=[r.q_liq for r in res], y=mid, width=h, orientation="h", marker_color="#9ec5f0", name="Qж"), 1, 5)
    fig.add_trace(go.Bar(x=[r.q_oil for r in res], y=mid, width=h, orientation="h", marker_color="#2e9d4f", name="Qн"), 1, 5)
    fig.add_trace(go.Bar(x=[r.fw if r.used else 0 for r in res], y=mid, width=h, orientation="h",
                         marker_color=["#d64545" if r.fw > comp.max_wc else "#7fb46a" for r in res], showlegend=False,
                         hovertext=[f"Fw {r.fw:.1f}%" for r in res], hoverinfo="text"), 1, 6)
    r = case.rules
    if r.pgmc_bottom > sl[0].top:
        pg_top = r.pgmc_top if 0 < r.pgmc_top < r.pgmc_bottom else sl[0].top
        fig.add_trace(go.Bar(x=[1], y=[(pg_top + r.pgmc_bottom) / 2], width=[r.pgmc_bottom - pg_top], orientation="h",
                             marker_color="#5b6abf", hovertext=f"ПГМЦ до {r.pgmc_bottom:.1f} м", hoverinfo="text",
                             showlegend=False), 1, 7)
    for e in liner(comp, case.rules):
        fig.add_trace(go.Bar(x=[1], y=[(e.top + e.bottom) / 2], width=[e.length], orientation="h",
                             marker_color=LINER_COLORS[e.kind], hovertext=f"{e.kind}: {e.top:.1f}–{e.bottom:.1f} м ({e.length:.1f} м)",
                             hoverinfo="text", showlegend=False), 1, 7)
    fig.update_layout(barmode="overlay", height=820, margin=dict(l=10, r=10, t=50, b=10), title=title,
                      legend=dict(orientation="h", y=-0.02), bargap=0)
    fig.update_yaxes(autorange="reversed", title_text="Глубина по стволу, м", row=1, col=1)
    fig.update_xaxes(type="log", row=1, col=2)
    fig.update_xaxes(type="log", row=1, col=4)
    fig.update_xaxes(range=[0, 1], row=1, col=3)
    fig.update_xaxes(range=[0, 100], row=1, col=6)
    for col in (1, 7):
        fig.update_xaxes(showticklabels=False, range=[0, 1], row=1, col=col)
    return fig


def _lateral_block(case, comp, mode_txt, key="auto"):
    """Планшет и таблицы одного ствола (key — различает подобранную и свою компоновку на странице)."""
    well = case.slices[0].well if case.slices else ""
    title = (f"{well}: своя компоновка" if key.startswith("manual") else
             f"{well}: компоновка при пределе обводнённости {mode_txt} {comp.max_wc:g}%")
    st.plotly_chart(_planshet(case, comp, title), use_container_width=True, key=f"planshet_{key}")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("**Фильтры**")
        st.dataframe(pd.DataFrame([{"Кровля, м": round(f.top, 1), "Подошва, м": round(f.bottom, 1),
                                    "Длина, м": round(f.length, 1), "Qж": round(f.q_liq, 1), "Qн": round(f.q_oil, 1),
                                    "Обв., %": round(f.water_cut, 1)} for f in comp.filters]), hide_index=True,
                     use_container_width=True, key=f"filters_{key}")
        st.markdown("**Раскладка хвостовика**")
        st.dataframe(pd.DataFrame([{"Элемент": e.kind, "Кровля, м": round(e.top, 1), "Подошва, м": round(e.bottom, 1),
                                    "Длина, м": round(e.length, 1)} for e in liner(comp, case.rules)]),
                     hide_index=True, use_container_width=True, height=260, key=f"liner_{key}")
    with c2:
        st.markdown("**Изоляция** — участки, которые не вскрываются")
        st.dataframe(pd.DataFrame([{"Кровля, м": round(i.top, 1), "Подошва, м": round(i.bottom, 1),
                                    "Причина": i.reason, "Потенциал Qж": round(i.q_liq, 1), "Потенциал Qн": round(i.q_oil, 1),
                                    "Обв., %": round(i.water_cut, 0),
                                    "Переток по пласту, %": round(i.cross_frac * 100),
                                    "Пакеры": "; ".join(f"{a:.1f}–{b:.1f}" for a, b in i.packers)} for i in comp.isolations]),
                     hide_index=True, use_container_width=True, key=f"iso_{key}")


def _watering_map(c_case, comp, fc, li: int, rho: float, title: str):
    """Тепловая карта: обводнённость каждого среза ствола по месяцам, линии — границы фильтров."""
    ks = [k for k in range(len(c_case.slices)) if (li, k) in fc.slice_q and fc.slice_q[(li, k)]]
    if not ks:
        return None
    n_m = max(len(fc.slice_q[(li, k)]) for k in ks)
    y = [(c_case.slices[k].top + c_case.slices[k].bottom) / 2 for k in ks]
    z = []
    for k in ks:
        row = []
        for t in fc.slice_q[(li, k)]:
            qo, ql = t[0], t[1]
            fw = 0.0 if ql <= 0 else max(0.0, min(100.0, 100 * (1 - qo / rho / ql)))
            row.append(fw)
        row += [None] * (n_m - len(row))
        z.append(row)
    fig = go.Figure(go.Heatmap(x=list(range(1, n_m + 1)), y=y, z=z, zmin=0, zmax=100, colorscale="RdYlGn_r",
                               colorbar=dict(title="обв., %"),
                               hovertemplate="месяц %{x}<br>глубина %{y:.1f} м<br>обводнённость %{z:.0f}%<extra></extra>"))
    for f in comp.filters:
        for yy in (f.top, f.bottom):
            fig.add_hline(y=yy, line=dict(color="#1f5f2e", width=1.5))
    fig.update_yaxes(autorange="reversed", title_text="глубина по стволу, м")
    fig.update_xaxes(title_text="месяц")
    fig.update_layout(title=title, height=460, margin=dict(l=10, r=60, t=40, b=10))
    return fig


def _cone_section(c_case, comp, fc, li: int, month: int, owc, goc, rules, title: str, barrier=None):
    """Разрез вдоль ствола: траектория по абс. отметке, ВНК/ГНК, поднявшаяся вода и опустившийся газ."""
    sl = c_case.slices
    pts = [(k, (s.top + s.bottom) / 2, (s.top_tvd + s.bottom_tvd) / 2) for k, s in enumerate(sl) if s.bottom_tvd > 0]
    if len(pts) < 2:
        return None
    md = [p[1] for p in pts]
    tvd = [p[2] for p in pts]
    fig = go.Figure()
    i = month - 1
    if owc:
        front = []
        for k, x_md, z in pts:
            d = fc.slice_info.get((li, k), {})
            fw = d.get("front_w") or []
            frac = fw[min(i, len(fw) - 1)] if fw and d.get("vertical") else 0.0
            front.append(owc - (d.get("d_owc") or 0.0) * frac)
        fig.add_trace(go.Scatter(x=md, y=[owc] * len(md), mode="lines", line=dict(color="#2f7fd8", width=1),
                                 name="ВНК", hoverinfo="skip"))
        fig.add_trace(go.Scatter(x=md, y=front, mode="lines", fill="tonexty", fillcolor="rgba(47,127,216,0.35)",
                                 line=dict(color="#2f7fd8", width=1.5), name="вода (фронт)",
                                 hovertemplate="%{x:.0f} м<br>фронт воды %{y:.1f} м АО<extra></extra>"))
    if goc:
        front = []
        for k, x_md, z in pts:
            d = fc.slice_info.get((li, k), {})
            fg = d.get("front_g") or []
            frac = fg[min(i, len(fg) - 1)] if fg and d.get("d_goc") is not None else 0.0
            front.append(goc + (d.get("d_goc") or 0.0) * frac)
        fig.add_trace(go.Scatter(x=md, y=[goc] * len(md), mode="lines", line=dict(color="#d64545", width=1),
                                 name="ГНК", hoverinfo="skip"))
        fig.add_trace(go.Scatter(x=md, y=front, mode="lines", fill="tonexty", fillcolor="rgba(214,69,69,0.30)",
                                 line=dict(color="#d64545", width=1.5), name="газ (фронт)",
                                 hovertemplate="%{x:.0f} м<br>фронт газа %{y:.1f} м АО<extra></extra>"))
    if barrier and barrier.get("top_tvd"):
        b_top, b_bot = barrier["top_tvd"], barrier["top_tvd"] + barrier.get("thickness", 0.0)
        fig.add_trace(go.Scatter(x=[md[0], md[-1], md[-1], md[0], md[0]], y=[b_top, b_top, b_bot, b_bot, b_top],
                                 fill="toself", fillcolor="rgba(120,120,120,0.45)", line=dict(color="#777", width=1),
                                 name=f"барьер ({barrier.get('perm', 0):g} мД)", hoverinfo="name"))
    fig.add_trace(go.Scatter(x=md, y=tvd, mode="lines", line=dict(color="#555", width=2), name="ствол",
                             hovertemplate="%{x:.0f} м по стволу<br>%{y:.1f} м АО<extra></extra>"))
    for f in comp.filters:
        seg = [(m_, z) for m_, z in zip(md, tvd) if f.top <= m_ <= f.bottom]
        if seg:
            fig.add_trace(go.Scatter(x=[a for a, _ in seg], y=[b for _, b in seg], mode="lines",
                                     line=dict(color="#2e9d4f", width=7), name="фильтр", showlegend=f is comp.filters[0],
                                     hoverinfo="skip"))
    pk = [(a + b) / 2 for iso in comp.isolations for a, b in iso.packers]
    if pk:
        import numpy as np
        fig.add_trace(go.Scatter(x=pk, y=list(np.interp(pk, md, tvd)), mode="markers", name="пакер",
                                 marker=dict(symbol="square", size=9, color="#e07b00"), hoverinfo="skip"))
    fig.update_yaxes(autorange="reversed", title_text="абс. отметка, м")
    fig.update_xaxes(title_text="глубина по стволу, м")
    fig.update_layout(title=title, height=420, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.2))
    return fig


def _dynamics_charts(cases, variants, fkw, rho, rules, state):
    """Раздел «Обводнение ствола и конусы»."""
    st.markdown("**Обводнение ствола во времени и конусы**")
    names = [v[0] for v in variants]
    c1, c2 = st.columns([2, 3])
    pick = c1.selectbox("Вариант", names, key="dyn_variant")
    name, w, f, _col = variants[names.index(pick)]
    if not f.months:
        st.info("Для этого варианта прогноза нет.")
        return
    li = 0
    if len(cases) > 1:
        li = c2.selectbox("Ствол", list(range(len(cases))), format_func=lambda i: w.names[i], key="dyn_lateral")
    fig = _watering_map(cases[li], w.laterals[li], f, li, rho,
                        f"Обводнённость срезов по месяцам — {name}")
    if fig is not None:
        st.plotly_chart(fig, use_container_width=True, key=f"wmap_{name}_{li}")
        st.caption("Каждая строка — срез РИГИС, цвет — его обводнённость в этом месяце (у изолированных — то, что "
                   "они отдают перетоком). Тёмно-зелёные линии — границы фильтров. Видно, где и когда к стволу "
                   "приходит вода.")
    owc, goc = fkw.get("owc_tvd"), fkw.get("goc_tvd")
    if fkw.get("displacement") == "bl" and (owc or goc):
        n = len([m for m in f.months if m.working]) or 1
        month = st.slider("Месяц разреза", 1, n, n, key=f"dyn_month_{name}")
        fig = _cone_section(cases[li], w.laterals[li], f, li, month, owc, goc, rules,
                            f"Разрез вдоль ствола, месяц {month} — {name}", fkw.get("barrier"))
        if fig is not None:
            st.plotly_chart(fig, use_container_width=True, key=f"section_{name}_{li}")
            st.caption("Вода поднимается от ВНК, газ опускается от ГНК (Баклей–Леверетт с гравитацией). Когда фронт "
                       "доходит до ствола — прорыв. Показана только подошвенная часть; вода сбоку (законтурная, ППД) "
                       "на разрезе не видна — её прорыв в таблице ниже («Откуда вода»). Масштаб по вертикали растянут.")
    elif fkw.get("displacement") == "bl":
        st.caption("Разрез с конусами появится, если задать ВНК или ГНК в «Контактах» (абс. отметка).")


def _forecast_settings(case):
    """Параметры прогноза, перетоков и экономики (нужны и подбору на накопленную добычу / NPV)."""
    with st.expander("Прогноз, перетоки из изолированных зон, экономика", expanded=False):
        st.markdown("**Прогноз**")
        f1, f2, f3, f4, f5, f6 = st.columns(6)
        years = f1.number_input("Горизонт, лет", 1, 30, 5)
        vrr = f2.number_input("Компенсация отбора", 0.0, 1.5, 0.8, 0.05, help="0 — истощение, 1 — давление держится")
        k_oper = f3.number_input("Кэ", 0.1, 1.0, 0.95, 0.01)
        q_min = f4.number_input("Мин. Qн, т/сут", 0.0, 500.0, 5.0, 1.0)
        wc_max = f5.number_input("Предельная обв., %", 1.0, 100.0, 98.0, 1.0)
        p_wf = f6.number_input("Рзаб в прогнозе, атм", 1.0, 500.0, float(case.defaults.p_wf), 1.0)
        disp = st.radio("Вытеснение", ["Баклей–Леверетт: фронт, конусы от ВНК и ГНК",
                                       "упрощённое: обводнение сразу (как раньше)"], horizontal=True,
                        help="Баклей–Леверетт: вода (и газ шапки) приходит фронтом; при заданных в «Контактах» "
                             "ВНК/ГНК — снизу и сверху через столб пласта с гравитацией (конусы).")
        w1, w2 = st.columns([3, 2])
        src = w1.radio("Источник воды (для Баклея–Леверетта)",
                       ["смешанный: снизу и сбоку", "снизу — подошвенная вода от ВНК",
                        "сбоку — законтурная вода, ППД"], horizontal=True,
                       help="Снизу — конус от ВНК (нужен ВНК в «Контактах»); сбоку — фронт с латерали "
                            "(законтурная вода, нагнетательные скважины). Смешанный — компенсация делится по доле.")
        share = 1.0 if src.startswith("снизу") else 0.0 if src.startswith("сбоку") else w2.slider(
            "Доля подошвенной воды", 0.0, 1.0, 0.5, 0.05,
            help="Какая часть компенсации приходит снизу от ВНК; остальное — сбоку")
        st.markdown("**Перетоки через пакеры** (негерметичность, износ эластомеров)")
        l1, l2, _ = st.columns([1, 1, 4])
        leak = l1.number_input("Доля потенциала, %", 0.0, 100.0, 0.0, 1.0,
                               help="Какую долю своего дебита изолированный коллектор отдаёт в фильтры; 0 — пакеры герметичны")
        growth = l2.number_input("Рост в год, п.п.", 0.0, 100.0, 0.0, 1.0, help="Износ эластомеров")
        st.markdown("**Экономика** (значения по умолчанию условные — задайте свои)")
        d = Economics()
        e1, e2, e3, e4 = st.columns(4)
        econ = Economics(
            oil_price=e1.number_input("Нетбэк нефти, руб./т", 0.0, 1e6, d.oil_price, 500.0),
            liquid_cost=e2.number_input("Подъём жидкости, руб./м³", 0.0, 1e5, d.liquid_cost, 10.0),
            water_cost=e3.number_input("Подготовка воды, руб./м³", 0.0, 1e5, d.water_cost, 10.0),
            discount=e4.number_input("Ставка дисконт., %/год", 0.0, 100.0, d.discount, 1.0),
            filter_m=e1.number_input("Фильтр, тыс. руб./м", 0.0, 1e4, d.filter_m, 1.0),
            blank_m=e2.number_input("Глухая труба, тыс. руб./м", 0.0, 1e4, d.blank_m, 1.0),
            packer=e3.number_input("Пакер-эластомер, тыс. руб./шт", 0.0, 1e5, d.packer, 10.0),
            fixed=e4.number_input("Прочее на компоновку, тыс. руб.", 0.0, 1e6, d.fixed, 100.0))
    fkw = dict(months=int(years * 12), vrr=vrr, k_oper=k_oper, q_min=q_min, wc_max=wc_max,
               leak=leak / 100, leak_growth=growth / 100, displacement="bl" if disp.startswith("Баклей") else "piston",
               bottom_share=share)
    return fkw, p_wf, econ, int(years)


def render():
    st.header("Подбор компоновки по РИГИС")
    st.caption("Загрузите РИГИС по скважине — инструмент подберёт интервалы фильтров на максимум Qн, накопленной "
               "нефти или NPV при пределе обводнённости и покажет, какие участки изолировать пакерами. Расчёт срезов — формулы листа «РИГИС» "
               "RE Tools (сверено с книгой), подбор — точный, вместо «Поиска решения». Несколько стволов (МЗС) "
               "определяются по столбцу «скважина» выгрузки.")
    _sources()
    state = st.session_state
    if "case" not in state:
        st.info("Выберите источник данных.")
        return
    _object_editor()
    case = _settings(state.case)
    obj, pal = state.obj, state.pal
    cases = split_laterals(case)
    if len(cases) > 1:
        st.info(f"МЗС: {len(cases)} ствола — " + ", ".join(f"«{c.slices[0].well}» ({len(c.slices)} срезов)" for c in cases)
                + ". Подбор совместный: общее число фильтров, общая длина секций, предел обводнённости.")
    fkw, p_wf, econ, years = _forecast_settings(case)
    zones = [(i.top, i.bottom) for i in case.intervals]
    fkw["zones"] = [zones] * len(cases)
    cf = state.get("crossflow", {"crossflow": True, "trans_mult": 1.0})
    contacts = state.get("contacts", {})
    fkw.update(cf, owc_tvd=contacts.get("owc_tvd"), goc_tvd=contacts.get("goc_tvd"), barrier=state.get("barrier"))
    if fkw["owc_tvd"] is None:
        fkw["bottom_share"] = 0.0          # без ВНК вода только сбоку
    fcase = [replace(c, defaults=replace(c.defaults, p_wf=p_wf)) for c in cases]
    st.subheader("3. Подбор")
    goal = st.radio("Цель подбора", list(OBJECTIVES), format_func=OBJECTIVES.get, horizontal=True,
                    help="Накопленная нефть и NPV — за горизонт прогноза с учётом обводнения срезов, падения давления, "
                         "перетоков и стоимости компоновки (настройки выше). Ограничения те же, что при подборе на Qн.")
    opts = dict(max_filters=state.get("max_filters", 10), zones=[zones] * len(cases),
                wc_mode=state.get("wc_mode", "filter"), gas_offset=state.get("gas_offset", 0.0),
                filter_cost=state.get("filter_cost", 0.0), **state.get("contacts", {}), **cf)
    key = (repr(case), id(obj), id(pal), repr(opts), goal, repr(fkw), p_wf, repr(econ))
    if st.button("Подобрать компоновку", type="primary") or "wcomp" not in state or state.get("comp_key") != key:
        with st.spinner("Подбор…" if goal == "rate" else "Подбор с прогнозом по итерациям…"):
            state.lp = plan_lifetime(cases, obj, pal, goal, econ, fkw, fcase, **opts)
            state.wcomp = state.lp.wcomp
            state.curve = wc_curve_multi(cases, obj, pal, **opts)
            ranges = []
            for c in cases:
                used = [(s.top, s.bottom) for s, r in zip(c.slices, compute(c, obj, pal, zones)) if r.used]
                ranges.append([(used[0][0], used[-1][1])] if used else [])
            state.all_open = evaluate_multi(cases, obj, pal, ranges, [zones] * len(cases))
            state.fc_open = forecast(fcase, obj, pal, state.all_open.laterals, state.all_open.l_total, **fkw)
            state.comp_key = key
    wc = state.wcomp
    ao = state.all_open
    lp = state.lp
    if goal != "rate":
        gain = lp.score - lp.base_score
        st.success(f"{OBJECTIVES[goal].capitalize()} за {years} лет: {lp.score:,.1f} {lp.unit} — на {gain:+,.1f} {lp.unit} "
                   f"к подбору на начальный Qн ({lp.base_score:,.1f}). Итераций прогноза: {len(lp.history)}."
                   .replace(",", " "))
    rho = obj.rho_oil
    qw = lambda w: max(w.q_liq - w.q_oil / rho, 0.0)   # noqa: E731  вода, м³/сут
    k = st.columns(6)
    k[0].metric("Qн, т/сут", f"{wc.q_oil:.1f}", f"{wc.q_oil - ao.q_oil:+.1f} к «всё открыто»")
    k[1].metric("Qж, м³/сут", f"{wc.q_liq:.1f}")
    k[2].metric("Qв, м³/сут", f"{qw(wc):.1f}", f"{qw(wc) - qw(ao):+.1f}", delta_color="inverse")
    k[3].metric("Обводнённость, %", f"{wc.water_cut:.1f}", f"{wc.water_cut - ao.water_cut:+.1f} п.п.", delta_color="inverse")
    k[4].metric("Фильтров", sum(len(c.filters) for c in wc.laterals))
    k[5].metric("Длина фильтров, м", f"{sum(f.length for c in wc.laterals for f in c.filters):.0f}")
    st.caption(f"Для сравнения, весь коллектор открыт: Qн {ao.q_oil:.1f} т/сут, Qж {ao.q_liq:.1f} м³/сут, "
               f"Qв {qw(ao):.1f} м³/сут, обв. {ao.water_cut:.1f}%."
               + (f" В итоге подобранной компоновки — переток по пласту из изолированных участков: "
                  f"Qж {wc.cross_liq:.1f} м³/сут, Qн {wc.cross_oil:.1f} т/сут." if cf["crossflow"] else ""))
    mode_txt = "фильтра" if state.get("wc_mode", "filter") == "filter" else "скважины"
    if len(cases) == 1:
        _lateral_block(cases[0], wc.laterals[0], mode_txt)
    else:
        st.dataframe(pd.DataFrame([{"Ствол": n, "Qн, т/сут": round(c.q_oil, 1), "Qж, м³/сут": round(c.q_liq, 1),
                                    "Обв., %": round(c.water_cut, 1), "Фильтров": len(c.filters)}
                                   for n, c in zip(wc.names, wc.laterals)]), hide_index=True)
        for j, (tab, c_case, comp) in enumerate(zip(st.tabs(wc.names), cases, wc.laterals)):
            with tab:
                _lateral_block(c_case, comp, mode_txt, key=f"auto_{j}")
    cv = pd.DataFrame([{"предел": c.max_wc, "Qн": c.q_oil, "обв": c.water_cut} for c in state.curve])
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=cv["предел"], y=cv["Qн"], mode="lines+markers", name="Qн", line=dict(color="#2e9d4f")))
    fig.add_trace(go.Scatter(x=[wc.max_wc], y=[wc.q_oil], mode="markers", marker=dict(size=13, color="#e07b00"),
                             name="текущий предел"))
    fig.update_layout(title=f"Qн от предела обводнённости {mode_txt}", height=320,
                      xaxis_title=f"предел обводнённости {mode_txt}, %", yaxis_title="Qн, т/сут",
                      margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.3))
    st.plotly_chart(fig, use_container_width=True)

    # --- своя компоновка ---
    st.subheader("Своя компоновка — интервалы вручную")
    st.caption("Впишите свои интервалы фильтров (по глубине ствола) — дебиты, изоляция, пакеры и раскладка хвостовика "
               "считаются по тем же правилам, что и у подбора. Фильтр вскрывает срезы РИГИС, середина которых "
               "попадает в интервал, поэтому границы фильтра — по границам срезов. Сначала в таблице — интервалы "
               "подобранной компоновки.")
    if st.button("Сбросить к подобранной компоновке"):
        state.manual_ver = state.get("manual_ver", 0) + 1
    man_ranges = []
    holders = [st.container()] if len(cases) == 1 else st.tabs(wc.names)
    for i, (holder, comp) in enumerate(zip(holders, wc.laterals)):
        with holder:
            init = pd.DataFrame([{"Кровля, м": round(f.top, 2), "Подошва, м": round(f.bottom, 2)} for f in comp.filters],
                                columns=["Кровля, м", "Подошва, м"])
            ed = st.data_editor(init, num_rows="dynamic", hide_index=True,
                                key=f"manual_{i}_{state.get('source', '')}_{state.get('manual_ver', 0)}")
            rr = [(float(r["Кровля, м"]), float(r["Подошва, м"])) for _, r in ed.iterrows()
                  if pd.notna(r["Кровля, м"]) and pd.notna(r["Подошва, м"]) and float(r["Подошва, м"]) > float(r["Кровля, м"])]
            man_ranges.append(sorted(rr))
    man, notes = manual_plan(cases, obj, pal, man_ranges, [zones] * len(cases), max_wc=wc.max_wc,
                             gas_offset=state.get("gas_offset", 0.0), names=wc.names, **state.get("contacts", {}),
                             **cf)
    n_pk = lambda w: sum(len(i.packers) for c in w.laterals for i in c.isolations)   # noqa: E731
    cmp = []
    for name, w in (("своя", man), ("подобранная", wc), ("всё открыто", ao)):
        cmp.append({"Вариант": name, "Qн, т/сут": round(w.q_oil, 1), "Qж, м³/сут": round(w.q_liq, 1),
                    "Qв, м³/сут": round(qw(w), 1), "Обв., %": round(w.water_cut, 1),
                    "Фильтров": sum(len(c.filters) for c in w.laterals),
                    "Длина фильтров, м": round(sum(f.length for c in w.laterals for f in c.filters)),
                    "Пакеров": n_pk(w) if name != "всё открыто" else 0})
    st.dataframe(pd.DataFrame(cmp), hide_index=True, use_container_width=True)
    for n in notes:
        st.warning(n)
    with st.expander("Планшет, фильтры и изоляция своей компоновки", expanded=False):
        if len(cases) == 1:
            _lateral_block(cases[0], man.laterals[0], mode_txt, key="manual_0")
        else:
            for j, (tab, c_case, comp) in enumerate(zip(st.tabs(man.names), cases, man.laterals)):
                with tab:
                    _lateral_block(c_case, comp, mode_txt, key=f"manual_{j}")
    man_key = (key, repr(man_ranges))
    if state.get("fc_man_key") != man_key:
        state.fc_man = forecast(fcase, obj, pal, man.laterals, man.l_total, **fkw)
        state.fc_man_key = man_key

    st.subheader("4. Прогноз добычи и экономика")
    st.caption("Помесячно на той же физике срезов: давление — по матбалансу объекта от отбора нефти и воды, "
               "компенсация замещает отобранный объём водой и держит давление, изолированные зоны отдают в фильтры "
               "заданную долю потенциала (перетоки). Параметры — в блоке «Прогноз, перетоки, экономика» выше.")
    fc, fo = lp.forecast, state.fc_open
    variants = [("подобранная", wc, fc, "#2e9d4f")]
    if goal != "rate":
        variants.append(("на начальный Qн", lp.base, lp.base_forecast, "#e07b00"))
    variants.append(("своя", man, state.fc_man, "#5b6abf"))
    variants.append(("всё открыто", ao, fo, "#9aa0a6"))
    fig = make_subplots(rows=1, cols=4, subplot_titles=("Qн, т/сут", "Обводнённость, %", "Рпл, атм",
                                                        "Накопленный NPV, млн руб."))
    for name, w, f, col in variants:
        t = [m.month for m in f.months]
        fig.add_trace(go.Scatter(x=t, y=[m.q_oil for m in f.months], name=name, line=dict(color=col)), 1, 1)
        fig.add_trace(go.Scatter(x=t, y=[m.water_cut for m in f.months], name=name, line=dict(color=col),
                                 showlegend=False), 1, 2)
        fig.add_trace(go.Scatter(x=t, y=[m.p_res for m in f.months], name=name, line=dict(color=col),
                                 showlegend=False), 1, 3)
        cf = cash_flow(f, econ)
        cap = capex(cases, w, econ).total
        fig.add_trace(go.Scatter(x=[0] + [r["month"] for r in cf], y=[-cap] + [r["npv"] - cap for r in cf], name=name,
                                 line=dict(color=col), showlegend=False), 1, 4)
    fig.update_xaxes(title_text="месяц")
    fig.update_layout(height=320, margin=dict(l=10, r=10, t=40, b=10), legend=dict(orientation="h", y=-0.3))
    st.plotly_chart(fig, use_container_width=True)

    def cum_at(f, month):
        ms = [m for m in f.months if m.month <= month]
        return (ms[-1].cum_oil, ms[-1].cum_liq) if ms else (0.0, 0.0)

    rows = []
    for yr in sorted({1, 3, years}):
        if yr > years:
            continue
        for name, _, f, _c in variants:
            o, l = cum_at(f, yr * 12)
            rows.append({"Лет": yr, "Вариант": name, "Нефть, тыс. т": round(o, 1), "Жидкость, тыс. м³": round(l, 1),
                         "Вода, тыс. м³": round(max(l - o / obj.rho_oil, 0), 1)})
    st.dataframe(pd.DataFrame(rows), hide_index=True)
    econ_rows = []
    for name, w, f, _c in variants:
        cap = capex(cases, w, econ)
        econ_rows.append({"Вариант": name, "Qн нач., т/сут": round(w.q_oil, 1), "Нефть, тыс. т": round(f.cum_oil, 1),
                          "Вода, тыс. м³": round(f.cum_water, 1), "Переток, тыс. м³": round(f.cum_leak_liq, 1),
                          "Фильтр, м": round(cap.filter_len), "Глухая труба, м": round(cap.blank_len),
                          "Пакеров": cap.packers, "Компоновка, млн руб.": round(cap.total, 1),
                          "NPV, млн руб.": round(npv(f, econ, cap.total), 1), "Остановка": f.stop_reason})
    _dynamics_charts(cases, variants, fkw, obj.rho_oil, case.rules, state)
    if fkw.get("displacement") == "bl":
        st.markdown("**Прорыв воды и газа по фильтрам** (подобранная компоновка, Баклей–Леверетт)")
        brows = []
        for li, (c_case, comp) in enumerate(zip(cases, wc.laterals)):
            for n_f, flt in enumerate(comp.filters, 1):
                ks = [k for k, s_ in enumerate(c_case.slices) if flt.top <= s_.top and s_.bottom <= flt.bottom]
                inf = [fc.slice_info.get((li, k), {}) for k in ks]
                inf = [d for d in inf if d]
                if not inf:
                    continue
                bw = [d["bt_w"] for d in inf if d.get("bt_w")]
                bg = [d["bt_g"] for d in inf if d.get("bt_g")]
                vert = [d for d in inf if d.get("vertical")]
                unstable = [d for d in vert if d.get("flux_w", 0) > d.get("q_crit_w", float("inf"))]
                first = min((d for d in inf if d.get("bt_w")), key=lambda d: d["bt_w"], default=None)
                brows.append({"Ствол": wc.names[li], "Фильтр": n_f, "Кровля, м": round(flt.top, 1),
                              "Подошва, м": round(flt.bottom, 1),
                              "Прорыв воды, мес": min(bw) if bw else "нет",
                              "Откуда вода": (first.get("src_w") or "—") if first else "—",
                              "Срезов с прорывом воды": len(bw),
                              "Прорыв газа, мес": min(bg) if bg else "нет",
                              "До ВНК, м (мин.)": round(min(d["d_owc"] for d in vert), 1) if vert else "—",
                              "Под барьером, срезов": sum(1 for d in vert if d.get("barrier")) if vert else "—",
                              "Фронт неустойчив (Дитц), срезов": f"{len(unstable)} из {len(vert)}" if vert else "—"})
        if brows:
            st.dataframe(pd.DataFrame(brows), hide_index=True, use_container_width=True)
        st.caption("Прорыв — первый месяц, когда фронт вытесняющей фазы дошёл до забоя среза. «Фронт неустойчив» — "
                   "приток воды к срезу выше критического дебита Дитца: конус растёт. ВНК и ГНК задаются в "
                   "«Контактах» (абс. отметка); без них — фронт с латерали.")
    st.markdown(f"**Итоги за {years} лет**")
    st.dataframe(pd.DataFrame(econ_rows), hide_index=True, use_container_width=True)
    st.caption(f"НГЗ дренируемого объёма подобранной компоновки: {fc.oip:.0f} тыс. т. Стоимость компоновки — по "
               "раскладке хвостовика на интервале от первого до последнего работающего среза.")
    state.econ_rows = econ_rows

    st.subheader("5. Выгрузка")
    c1, c2 = st.columns(2)
    c1.download_button("Компоновка и расчёт в Excel", well_xlsx(cases, wc, state.curve, title=f"Подбор: {state.get('source', '')}",
                                                               forecasts={n: f for n, _, f, _c in variants},
                                                               economics=econ_rows),
                       file_name="компоновка_РИГИС.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    c1.download_button("Своя компоновка в Excel", well_xlsx(cases, man, None, title=f"Своя компоновка: {state.get('source', '')}",
                                                            forecasts={"своя": state.fc_man}),
                       file_name="своя_компоновка_РИГИС.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    tmp = Path(tempfile.gettempdir()) / "retools_project.json"
    save_project(tmp, case, obj, pal)
    c2.download_button("Проект (JSON)", tmp.read_bytes(), file_name="проект_РИГИС.json", mime="application/json")
